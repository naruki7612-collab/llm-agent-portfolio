"""8/8版（旧版/update_zaigen_v9.py）と現行を同条件で走らせ、千円値の丸め廃止による差を見る。

確認したいこと
  ① 現行は原本と同じ「検収額(円)/1000 の小数」を保存する（旧版は整数に丸めていた）
  ② 旧版は round() 比較で「増えているのに黄色」の過剰検知を起こす。現行は起こさない
  ③ 書込セル数・スキップ挙動・安全停止は旧版と同じ（丸め以外は変えていない）
"""

import csv
import io
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as CL

HERE = Path(__file__).parent
# 8/8版は月ラベルが日付型でないと読めないため、旧方式の雛形を使う
XLSX = HERE / "旧版" / "雛形方式" / "財源集計表_正式版.xlsx"
CSV = HERE / "入力_実物" / "月度リース割賦情報管理表.csv"   # 8月のみの部分集合（当時の検証条件）
YELLOW = "00FFFF00"

EXCLUDE = {"競合負け", "自主謝絶", "否決", "没"}
_STRIP = str.maketrans("", "", " 　\t\r\n-_./\\・")


def norm_branch(name: str) -> str:
    s = unicodedata.normalize("NFKC", str(name)).translate(_STRIP)
    s = re.sub(r"(支店|チーム|営業)", "", s).lower().strip()
    return {"本社部": "本社店口"}.get(s, s)


def expected_sen() -> dict:
    """CSV から「部店 → ABC/D の千円値（丸めなし）」を作る。これが原本の書き方。"""
    b = CSV.read_bytes()
    if b.startswith(b"\xef\xbb\xbf"):
        b = b[3:]
    agg = defaultdict(float)
    for r in csv.DictReader(io.StringIO(b.decode("cp932"))):
        if r["進捗"] in EXCLUDE:
            continue
        k = r["確度"].strip().upper()
        kind = "ABC" if k in ("A", "B", "C") else ("D" if k == "D" else None)
        if not kind:
            continue
        agg[(norm_branch(r["部店"]), kind)] += float(r["検収額"] or 0)
    return {k: round(v / 1000, 3) for k, v in agg.items()}


def run(script: Path, run_date: str):
    work = Path(tempfile.mkdtemp(prefix="marume_"))
    (work / "uploads").mkdir()
    (work / "output").mkdir()
    shutil.copy(XLSX, work / "uploads" / XLSX.name)
    shutil.copy(CSV, work / "uploads" / CSV.name)
    lines = script.read_text().split("\n")
    # モジュールdocstringの終端直後に入れる（行番号を決め打ちしない）
    end = next(i for i in range(1, len(lines)) if lines[i].rstrip() == '"""')
    lines.insert(end + 1, "from __future__ import annotations")   # ローカル3.9対策のみ
    local = work / "run_local.py"
    local.write_text("\n".join(lines))
    p = subprocess.run([sys.executable, str(local), run_date],
                       cwd=work, capture_output=True, text=True)
    return p.stdout + p.stderr, work / "output" / "財源状況集計表_更新版.xlsx"


BRANCH_ROWS = {"本社1": 5, "本社2": 7, "輸送機": 9, "本社店口": 11, "北九州1": 15,
               "北九州2": 17, "久留米": 21, "熊本": 23, "大分": 25, "長崎": 27,
               "東京": 29, "大阪": 31, "kli不動産": 33}


def colored_rows(xlsx: Path, col: int) -> set:
    ws = openpyxl.load_workbook(xlsx)["集計表"]
    return {r for r in range(5, 37)
            if ws.cell(r, col).fill and ws.cell(r, col).fill.start_color.rgb == YELLOW}


def main() -> int:
    old, cur = HERE / "旧版" / "update_zaigen_v9.py", HERE / "update_zaigen.py"
    exp = expected_sen()
    ok = True

    # ---------- ① 保存値の型 ----------
    print("=" * 74)
    print("① 保存値：原本と同じ小数になるか（実行日 2026-08-20 / R列）")
    print("=" * 74)
    log_o, out_o = run(old, "2026-08-20")
    log_c, out_c = run(cur, "2026-08-20")
    ws_o = openpyxl.load_workbook(out_o)["集計表"]
    ws_c = openpyxl.load_workbook(out_c)["集計表"]
    col = 18  # R = 8/20
    print(f"{'部店':<12}{'期待(円/1000)':>14}{'旧版':>10}{'現行':>14}  判定")
    miss_o = miss_c = 0
    for b, r in BRANCH_ROWS.items():
        e = exp.get((b, "ABC"), 0.0)
        a, c = ws_o.cell(r, col).value, ws_c.cell(r, col).value
        if a != e:
            miss_o += 1
        if c != e:
            miss_c += 1
        print(f"{b:<12}{e:>14.3f}{a:>10}{c:>14}  {'一致' if c == e else '不一致'}")
    print()
    print(f"  旧版 が期待値と違うセル: {miss_o} / {len(BRANCH_ROWS)}")
    print(f"  現行 が期待値と違うセル: {miss_c} / {len(BRANCH_ROWS)}")
    ok &= miss_c == 0
    print(f"  [{'OK ' if miss_c == 0 else 'NG '}] 現行は原本と同じ小数を保存する")

    # ---------- ② 色付けの漏れ ----------
    print()
    print("=" * 74)
    print("② 色付け：前日と同額（前日以下）の部店に色が付くか（実行日 2026-08-06）")
    print("=" * 74)
    log_ob, out_ob = run(old, "2026-08-06")
    log_cb, out_cb = run(cur, "2026-08-06")
    col2 = 9  # I = 8/6
    c_o, c_c = colored_rows(out_ob, col2), colored_rows(out_cb, col2)
    row2b = {v: k for k, v in BRANCH_ROWS.items()}
    src = openpyxl.load_workbook(XLSX)["集計表"]
    print(f"{'部店':<12}{'前日8/5':>10}{'当日(真値)':>12}  {'旧版':>4}{'現行':>5}  判定")
    for b, r in BRANCH_ROWS.items():
        prev = src.cell(r, 8).value
        now = exp.get((b, "ABC"), 0.0)
        if prev is None:
            continue
        should = now <= prev + 1e-9
        m_o = "黄" if r in c_o else "−"
        m_c = "黄" if r in c_c else "−"
        tag = ""
        if (r in c_c) != should:
            tag = "★現行が仕様と不一致"
            ok = False
        elif (r in c_o) != should:
            tag = "←旧版は誤判定（増えているのに黄）" if r in c_o else "←旧版で漏れ"
        print(f"{b:<12}{prev:>10.3f}{now:>12.3f}  {m_o:>4}{m_c:>5}  {tag}")
    print()
    kajo = sorted(row2b.get(r, r) for r in (c_o - c_c))    # 旧版が余計に付けた
    more = sorted(row2b.get(r, r) for r in (c_c - c_o))    # 現行で新たに付いた
    print(f"  旧版の色付け: {len(c_o)}セル / 現行: {len(c_c)}セル")
    print(f"  旧版が誤って色を付けていた部店（増加なのに黄）: {kajo}")
    print(f"  現行で新たに色が付いた部店: {more}")
    print()
    print("  ※ round() は単調なので today<=prev なら round(today)<=round(prev) が必ず成立する。")
    print("     つまり旧版の丸め比較は取りこぼし（false negative）は起こさず、")
    print("     増えているのに黄色にする過剰検知（false positive）だけを起こす。")
    ok &= all((r in c_c) == (exp.get((row2b[r], "ABC"), 0.0) <= src.cell(r, 8).value + 1e-9)
              for r in BRANCH_ROWS.values() if src.cell(r, 8).value is not None)

    # ---------- ③ 丸め以外の挙動が同じか ----------
    print()
    print("=" * 74)
    print("③ 丸め以外の挙動：書込セル数・スキップ・安全停止")
    print("=" * 74)
    for label, l_o, l_c in [("2026-08-20", log_o, log_c), ("2026-08-06", log_ob, log_cb)]:
        for key in ["[STEP4] 書込セル数: 32 (想定 32)",
                    "表 2026-09: シート内に該当日付の列が見つかりません"]:
            a, c = key in l_o, key in l_c
            same = a == c
            ok &= same
            print(f"  [{'OK ' if same else 'NG '}] {label} 「{key[:38]}…」 旧版={a} 現行={c}")

    s_o, _ = run(old, "2026-08-11")
    s_c, o_c = run(cur, "2026-08-11")
    k = "ヘッダに実行日 2026-08-11 の列が見つかりません"
    same = (k in s_o) == (k in s_c)
    ok &= same and not o_c.exists()
    print(f"  [{'OK ' if same else 'NG '}] 2026-08-11 の安全停止が同じ / 現行は出力なし={not o_c.exists()}")

    print()
    print("=" * 74)
    print("総合判定:", "全項目 OK" if ok else "NG あり")
    print("=" * 74)
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
