"""翌月ぶんの財源集計表（骨組み）を作る。

ファイルは月ごとに新しく切り替える運用（先方確定・2026-09-08）。
毎月「ラベルを繰り上げ、日付列をその月の営業日に入れ替える」作業が発生するので、
その機械的な部分だけをこのツールで作る。

## 作るもの

    1つ目 = 当月 / 2つ目 = 翌月 / 3つ目 = 翌々月  のラベル（'9月' 形式）
    日付列 = 当月の営業日（土日祝を除く）。**3表とも同じ日付列**を共有する
    数式    = B列（要増加）・AB/AC/AD列（比率）は**原本のまま一切触らない**
    財源保有状況シート = 月の見出しだけを繰り上げ（原本にある列だけ）。
                        データ部は空にする

## 作らないもの（人が入れる）

    C列 目標   … 経営が決める数字。CSV にもシートのどこにも根拠が無い
    E列 月初   … 月初時点のスナップショット
    B列 要増加（1つ目の表だけ）… 2つ目・3つ目は数式で出る
    財源保有状況のデータ部（A/BC/D の実数）… 月初のスナップショット

## 触らないもの

    AB/AC/AD列（要増加比・月初比・前日比）の数式
        お客様が日々「最後に値を入れた列」に付け替えて運用している欄。
        こちらで参照先を動かすと運用と食い違うので原本のまま残す。
    財源保有状況で原本に見出しが無い列（V列）
        レポート部（B〜P）は T列と U列しか参照していない。
        見出しを足すと誰も読まないデータを書き続けることになる。

**目標は繰り上げてコピーしない。** 実データで確認したところ、翌月のうちは
据え置きだが「当月になる時点で改定」されていた
（7月版2つ目の8月目標 535 → 8月版1つ目の8月目標 700）。
勝手に前月の値を入れると、改定前の数字が入ったまま運用が進む危険がある。
空にしておけば update_zaigen.py 側の警告が名指しで拾う。

使い方:
    python make_month_file.py 2026-09
    python make_month_file.py 2026-09 --base 入力_実物/財源集計表_v1.xlsx
"""

import argparse
import datetime
import re
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as CL

HERE = Path(__file__).parent
DEFAULT_BASE = HERE / "入力_実物" / "財源集計表_v1.xlsx"

SHEET = "集計表"
HOLDING_SHEET = "財源保有状況"
MONTH_MARK, MONTH_MARK_COL = "月初", 5
DATE_FROM, DATE_TO = 6, 27          # F〜AA
TARGET_COL, GETSUSHO_COL = 3, 5     # C=目標, E=月初
RATIO_COLS = (28, 29, 30)           # AB=要増加比, AC=月初比, AD=前日比
HOLD_DATE_ROW, HOLD_FIRST_COL = 1, 20   # 財源保有状況: 1行目に月, T列から

# 国民の祝日（振替休日を含む）。年を足すときはここに追記する。
HOLIDAYS = {
    (2026, 1, 1), (2026, 1, 12), (2026, 2, 11), (2026, 2, 23), (2026, 3, 20),
    (2026, 4, 29), (2026, 5, 3), (2026, 5, 4), (2026, 5, 5), (2026, 5, 6),
    (2026, 7, 20), (2026, 8, 11), (2026, 9, 21), (2026, 9, 22), (2026, 9, 23),
    (2026, 10, 12), (2026, 11, 3), (2026, 11, 23),
    (2027, 1, 1), (2027, 1, 11), (2027, 2, 11), (2027, 2, 23), (2027, 3, 21),
    (2027, 3, 22), (2027, 4, 29), (2027, 5, 3), (2027, 5, 4), (2027, 5, 5),
    (2027, 7, 19), (2027, 8, 11), (2027, 9, 20), (2027, 9, 23), (2027, 10, 11),
    (2027, 11, 3), (2027, 11, 23),
}

_CELLREF = re.compile(r"^=\s*\$?([A-Z]{1,2})\$?(\d+)\s*$")


def shift_month(year: int, month: int, n: int) -> tuple:
    m = month + n
    return year + (m - 1) // 12, (m - 1) % 12 + 1


def business_days(year: int, month: int) -> list:
    days, d = [], datetime.date(year, month, 1)
    while d.month == month:
        if d.weekday() < 5 and (d.year, d.month, d.day) not in HOLIDAYS:
            days.append(d)
        d += datetime.timedelta(days=1)
    return days


def find_header_rows(ws) -> list:
    return [r for r in range(2, ws.max_row + 1)
            if str(ws.cell(r, MONTH_MARK_COL).value).strip() == MONTH_MARK]


def repoint_ratio(value, to_col: str):
    """'=H5-B5' の日付列参照を to_col に差し替える。'=H5-E5' 等も同じ形。"""
    if not (isinstance(value, str) and value.startswith("=")):
        return value
    return re.sub(r"^=\s*[A-Z]{1,2}(\d+)", rf"={to_col}\1", value, count=1)


def main() -> int:
    ap = argparse.ArgumentParser(description="翌月ぶんの財源集計表の骨組みを作る")
    ap.add_argument("month", help="対象月 YYYY-MM（例 2026-09）")
    ap.add_argument("--base", default=str(DEFAULT_BASE),
                    help="土台にするファイル（前月の原本など）")
    ap.add_argument("--out", default=None)
    a = ap.parse_args()

    try:
        year, month = (int(x) for x in a.month.split("-"))
        datetime.date(year, month, 1)
    except Exception:
        print(f"[失敗] 対象月の形式が不正です: {a.month}（YYYY-MM で指定）")
        return 1

    base = Path(a.base)
    if not base.is_absolute():
        base = HERE / base
    if not base.exists():
        print(f"[失敗] 土台のファイルが見つかりません: {base}")
        return 1
    out = HERE / (a.out or f"財源集計表_{year}年{month}月.xlsx")

    days = business_days(year, month)
    if len(days) > DATE_TO - DATE_FROM + 1:
        print(f"[失敗] {a.month} の営業日 {len(days)} 日が列数を超えます")
        return 1

    shutil.copy(base, out)
    wb = openpyxl.load_workbook(out)
    ws = wb[SHEET]
    headers = find_header_rows(ws)

    print(f"[土台] {base.name}")
    print(f"[対象] {year}年{month}月")
    print(f"[構成] 表 {len(headers)} 個（ヘッダ行 {headers}）\n")

    # ---- 1. ラベルを 当月 / 翌月 / 翌々月 に ----
    print("■ 月ラベル")
    date_fmt = ws.cell(headers[0], DATE_FROM).number_format
    label_fmt = ws.cell(headers[0] - 1, 1).number_format
    for i, hdr in enumerate(headers):
        y, m = shift_month(year, month, i)
        cell = ws.cell(hdr - 1, 1)
        before = cell.value
        cell.value = f"{m}月"
        cell.number_format = label_fmt
        note = "当月" if i == 0 else ("翌月" if i == 1 else f"{i}か月後")
        print(f"   A{hdr - 1}: {before!r} → '{m}月'  ({y}-{m:02d} / {note})")

    # ---- 2. 日付列を当月の営業日に。2つ目以降は1つ目への参照 ----
    print(f"\n■ 日付列（{len(days)}営業日 / 3表とも共通）")
    for i, hdr in enumerate(headers):
        for j, c in enumerate(range(DATE_FROM, DATE_TO + 1)):
            cell = ws.cell(hdr, c)
            if j >= len(days):
                cell.value = None
                continue
            if i == 0:
                cell.value = datetime.datetime(days[j].year, days[j].month, days[j].day)
                cell.number_format = date_fmt
            else:
                # 原本と同じ作り。1つ目の日付を参照させて共通化する
                cell.value = f"={CL(c)}{headers[0]}"
    print(f"   {headers[0]}行: {days[0]:%-m/%-d} 〜 {days[-1]:%-m/%-d} を実日付で")
    for hdr in headers[1:]:
        print(f"   {hdr}行: '={CL(DATE_FROM)}{headers[0]}' 形式で1つ目を参照")
    hol = [f"{mm}/{dd}" for (yy, mm, dd) in sorted(HOLIDAYS) if (yy, mm) == (year, month)]
    print(f"   祝日除外: {hol or 'なし'}")

    # ---- 3. 人が入れる欄を空にする ----
    print("\n■ 人が入れる欄を空にする")
    cleared = {"目標": 0, "月初": 0, "実績": 0, "要増加": 0}
    for i, hdr in enumerate(headers):
        end = headers[i + 1] - 1 if i + 1 < len(headers) else ws.max_row + 1
        for r in range(hdr + 1, end):
            if ws.cell(r, 1).value is None and ws.cell(r, 4).value is None:
                continue
            is_branch = ws.cell(r, 1).value not in (None, "本社計", "北九州計", "合計")
            for c in (TARGET_COL, GETSUSHO_COL):
                v = ws.cell(r, c).value
                if v is not None and not (isinstance(v, str) and v.startswith("=")):
                    ws.cell(r, c).value = None
                    cleared["目標" if c == TARGET_COL else "月初"] += 1
            for c in range(DATE_FROM, DATE_TO + 1):
                if ws.cell(r, c).value is not None:
                    ws.cell(r, c).value = None
                    cleared["実績"] += 1
            # 1つ目の表の B列（要増加）は実数入力。2つ目以降は数式なので触らない
            if i == 0 and is_branch:
                v = ws.cell(r, 2).value
                if v is not None and not (isinstance(v, str) and v.startswith("=")):
                    ws.cell(r, 2).value = None
                    cleared["要増加"] += 1
    for k, v in cleared.items():
        print(f"   {k}: {v}セル")

    # ---- 4. AB/AC/AD（比率）は触らない ----
    print("\n■ 比率の数式（AB/AC/AD）")
    print(f"   原本のまま（例: AB{headers[0] + 1} = "
          f"{ws.cell(headers[0] + 1, 28).value}）")
    print("   ※ お客様が日々付け替えて運用されている欄なので触りません")

    # ---- 5. 財源保有状況シート ----
    if HOLDING_SHEET in wb.sheetnames:
        h = wb[HOLDING_SHEET]
        print(f"\n■ {HOLDING_SHEET} シート")

        # 見出しは「原本にある列」だけ繰り上げる。無い列（V列）には足さない。
        cols = [c for c in range(HOLD_FIRST_COL, h.max_column + 1)
                if isinstance(h.cell(HOLD_DATE_ROW, c).value, datetime.datetime)]
        for i, col in enumerate(cols):
            y0, m0 = shift_month(year, month, i)
            y1, m1 = shift_month(year, month, i + 1)
            h.cell(HOLD_DATE_ROW, col).value = datetime.datetime(y0, m0, 1)
            h.cell(HOLD_DATE_ROW + 1, col).value = datetime.datetime(y1, m1, 1)
            print(f"   {CL(col)}列: {y0}-{m0:02d}（次月 {y1}-{m1:02d}）")
        skipped = [CL(c) for c in range(HOLD_FIRST_COL, HOLD_FIRST_COL + 4)
                   if c not in cols and any(h.cell(r, c).value is not None
                                            for r in range(1, h.max_row + 1))]
        if skipped:
            print(f"   {', '.join(skipped)}列: 原本に月の見出しが無いので触りません")

        # データ部は空にする（0 ではない）。月初のスナップショットで人が入れる欄。
        n = 0
        for r in range(1, h.max_row + 1):
            if not h.cell(r, HOLD_FIRST_COL - 1).value:      # S列に区分がある行だけ
                continue
            for c in cols:
                if h.cell(r, c).value is not None:
                    h.cell(r, c).value = None
                    n += 1
        print(f"   データ部を空に: {n}セル（0では埋めない。保有率が0扱いになるため）")

    wb.save(out)
    print(f"\n[保存] {out.name}")

    first = headers[0] + 1
    last = headers[0] + 32
    print("\n" + "=" * 66)
    print("お客様に入力していただく欄")
    print("=" * 66)
    print(f"  1. C列 目標   … 3表それぞれ（1つ目は C{first}〜C{last} の部店行）")
    print(f"  2. E列 月初   … 3表それぞれ（{year}年{month}月1日時点の値）")
    print(f"  3. B列 要増加 … 1つ目の表だけ（B{first}〜B{last}）。2つ目3つ目は数式で出ます")
    print("  4. 財源保有状況シートの T〜U列（部店別の A/BC/D）… 月初のスナップショット")
    print("\n  ※ 目標は前月ファイルから引き継いでいません。当月になる時点で")
    print("     改定されている実績があるため（8月目標 535 → 700）、")
    print("     古い値が残るのを避けて空にしています。")
    print("  ※ 空のまま実行しても当日列への書き込みは動きますが、")
    print("     update_zaigen.py が名指しで警告を出します。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
