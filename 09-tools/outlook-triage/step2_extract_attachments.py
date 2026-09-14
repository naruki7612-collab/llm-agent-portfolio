import json
import subprocess
import os

def step3_download_files(attachment_urls: list):
    """
    bash_execute で S3URL からファイルをダウンロード
    """
    print("[Step 3] ファイルをダウンロード中...")
    
    downloaded_files = []
    
    for idx, info in enumerate(attachment_urls, 1):
        s3_url = info.get("s3_url")
        ext = info.get("ext", "bin")
        filename = info.get("filename", f"file_{idx}")
        
        if not s3_url:
            print(f"  - {filename}: ⚠️  S3URLなし（スキップ）")
            continue
        
        print(f"  - {filename}")
        
        dest_path = f"tmp/attachment_file_{idx}.{ext}"
        
        # curl または aws s3 cp を使用してダウンロード
        if s3_url.startswith("s3://"):
            cmd = f"aws s3 cp {s3_url} {dest_path}"
        elif s3_url.startswith("https://"):
            cmd = f"curl -s -L '{s3_url}' -o {dest_path}"
        else:
            cmd = f"cp {s3_url} {dest_path}"
        
        try:
            result = subprocess.run(cmd, shell=True, capture_output=True, text=True, timeout=60)
            
            if result.returncode == 0 and os.path.exists(dest_path):
                downloaded_files.append({
                    "filename": filename,
                    "ext": ext,
                    "local_path": dest_path,
                    "s3_url": s3_url
                })
                print(f"      ✅ ダウンロード完了")
            else:
                print(f"      ❌ ダウンロード失敗: {result.stderr}")
        
        except Exception as e:
            print(f"      ❌ エラー: {e}")
    
    return downloaded_files

def step4_extract_text(downloaded_files: list):
    """
    ダウンロードしたファイルからテキストを抽出
    """
    print("\n[Step 4] テキストを抽出中...")
    
    attachment_contents = []
    
    for info in downloaded_files:
        local_path = info.get("local_path")
        ext = info.get("ext", "").lower()
        filename = info.get("filename")
        
        print(f"  - {filename}")
        
        content = ""
        
        try:
            if ext == "txt":
                # プレーンテキスト
                with open(local_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
            
            elif ext == "pdf":
                # PDF
                try:
                    result = subprocess.run(
                        ["pdftotext", local_path, "-"],
                        capture_output=True,
                        text=True,
                        timeout=30
                    )
                    if result.returncode == 0:
                        content = result.stdout
                    else:
                        content = "[PDF抽出失敗: pdftotext が利用できません]"
                except:
                    content = "[PDF抽出失敗: pdftotext が見つかりません]"
            
            elif ext in ["docx", "doc"]:
                # Word - python-docx を試す
                try:
                    from docx import Document
                    doc = Document(local_path)
                    for para in doc.paragraphs:
                        content += para.text + "\n"
                except ImportError:
                    content = "[DOCX抽出失敗: python-docx が必要です]"
                except Exception as e:
                    content = f"[DOCX抽出失敗: {e}]"
            
            elif ext in ["xlsx", "xls"]:
                # Excel - openpyxl を試す
                try:
                    from openpyxl import load_workbook
                    wb = load_workbook(local_path, data_only=True)
                    for sheet in wb.sheetnames[:3]:
                        content += f"[シート: {sheet}]\n"
                        ws = wb[sheet]
                        for row in list(ws.iter_rows(values_only=True))[:20]:
                            content += " | ".join(str(cell) if cell else "" for cell in row) + "\n"
                except ImportError:
                    content = "[XLSX抽出失敗: openpyxl が必要です]"
                except Exception as e:
                    content = f"[XLSX抽出失敗: {e}]"
            
            else:
                content = f"[抽出不可: {ext} 形式は未対応です]"
        
        except Exception as e:
            content = f"[抽出エラー: {e}]"
        
        attachment_contents.append({
            "filename": filename,
            "extension": ext,
            "content": content[:2000] if content else ""
        })
        
        print(f"      ✅ 抽出完了 ({len(content)} 文字)")
        
        # ダウンロードファイルを削除
        try:
            os.remove(local_path)
        except:
            pass
    
    return attachment_contents

def main():
    """
    メイン処理
    
    1. attachment_urls.json を読み込む
    2. ファイルをダウンロード
    3. テキストを抽出
    4. attachment_contents.json に保存
    """
    print(f"=== ファイルダウンロード・テキスト抽出処理 ===\n")
    
    # attachment_urls.json を確認
    if not os.path.exists("tmp/attachment_urls.json"):
        print("❌ tmp/attachment_urls.json が見つかりません")
        print("先に prepare_attachments.py を実行してください")
        return
    
    with open("tmp/attachment_urls.json", "r", encoding="utf-8") as f:
        url_data = json.load(f)
    
    attachment_urls = url_data.get("attachments", [])
    
    if not attachment_urls:
        print("✅ ダウンロード対象ファイルなし")
        with open("tmp/attachment_contents.json", "w", encoding="utf-8") as f:
            json.dump({"attachments": [], "count": 0, "status": "no_files"}, f, ensure_ascii=False)
        return
    
    print(f"対象ファイル数: {len(attachment_urls)}\n")
    
    # Step3: ファイルダウンロード
    downloaded_files = step3_download_files(attachment_urls)
    
    if not downloaded_files:
        print("\n⚠️  ダウンロードしたファイルなし")
        with open("tmp/attachment_contents.json", "w", encoding="utf-8") as f:
            json.dump({"attachments": [], "count": 0, "status": "download_failed"}, f, ensure_ascii=False)
        return
    
    # Step4: テキスト抽出
    attachment_contents = step4_extract_text(downloaded_files)
    
    # 結果を JSON に保存
    output_data = {
        "attachments": attachment_contents,
        "count": len(attachment_contents),
        "status": "completed"
    }
    
    with open("tmp/attachment_contents.json", "w", encoding="utf-8") as f:
        json.dump(output_data, f, ensure_ascii=False, indent=2)
    
    print(f"\n✅ 完了")
    print(f"抽出結果を保存しました: tmp/attachment_contents.json")
    print(f"\n次のコマンドでドラフト作成を実行してください:")
    print(f"  python step3_create_draft.py")

if __name__ == "__main__":
    main()
