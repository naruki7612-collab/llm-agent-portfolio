"""
epsilon_product_search.py — EPSILON商材検索 関数ライブラリ

メインAIが処理フローに沿って呼ぶ小さい関数集。オーケストレーション（カテゴリの
判断・商材の絞り込み）はメインAIが途中結果を見ながら行い、このライブラリは
決定的な処理（Excel読解・DB絞り込み・PDF名決定・取り寄せ・全文読解）だけを担う。

使い方（code_execute内。判断はステップ間でメインAIが行う）:
    cats = await get_categories()                 # ① カテゴリ一覧 {カテゴリ: 商材数}
    prods = await get_products("予約管理")         # ② カテゴリ内の商材（名前・業種・概要）
    pdfs = await get_pdfs("EPSILON SMART RESERVE",    # ③ 読むべきPDF名（用途別・コード決定）
                          "仕様・価格・その他")
    for a in await read_docs("<質問>", pdfs):      # ④ DL＆Gemini全文読解
        print(a["file"], a.get("answer") or a.get("error"))

診断ログ: tmp/debug_search.jsonl に全段を記録
"""

from __future__ import annotations

import asyncio
import copy
import datetime
import json
import logging
import os
import re
import time
import unicodedata

logger = logging.getLogger(__name__)

# ============================================================
# 定数
# ============================================================

INDEX_XLSX_NAME = "商材メタデータDB_v1.xlsx"  # 学習データ上の索引Excel名
# 索引ExcelのUUID（保険）。一覧照合で名前解決できるため通常は使わない。
INDEX_XLSX_UUID = "<INDEX_XLSX_UUID>"
_LIST_PAGE = 20  # rag_file_list のページサイズ（この値での全件取得はルームで実証済み）

# PDF読解はGemini明示時のみ動作するため読取りモデルを固定する（v30と同方針）
GEMINI_PDF_MODEL = "gemini/gemini-3.5-flash"

MAX_READ_FILES = 3                # 1質問で読むPDFの上限（読みすぎ防止）
MAX_PDF_BYTES = 64 * 1024 * 1024  # llm_call の添付サイズ上限（実測値）


class PdfTooLargeError(Exception):
    """llm_call添付がinline上限超過。決定論的エラーであり同じPDFの再送リトライは無意味。"""
    pass


_index_cache: list[dict] | None = None        # 索引24行のキャッシュ
_listing_cache: dict[str, str] | None = None  # _match_key(名前) → uuid
_uuid_cache: dict[str, str] = {}              # 解決済み name→uuid
_path_cache: dict[str, str] = {}              # DL済み name→sandbox_path

# ============================================================
# デバッグロガー（tmp/debug_search.jsonl に追記）— v30と同方式・截断なし
# ============================================================

_DEBUG_PATH = "tmp/debug_search.jsonl"
_START_TIME = time.time()
_SEQ = 0


def _describe(v, _depth=0):
    """任意の値を「型・長さ・中身」付きで完全記述する（截断なし）"""
    try:
        if v is None or isinstance(v, (bool, int, float)):
            return {"type": type(v).__name__, "value": v}
        if isinstance(v, str):
            return {"type": "str", "len": len(v), "value": v}
        if isinstance(v, dict):
            if _depth > 20:
                return {"type": "dict", "keys": list(v.keys()), "note": "max_depth_20"}
            return {"type": "dict", "keys": list(v.keys()),
                    "items": {str(k): _describe(val, _depth + 1) for k, val in v.items()}}
        if isinstance(v, (list, tuple)):
            if _depth > 20:
                return {"type": type(v).__name__, "len": len(v), "note": "max_depth_20"}
            return {"type": type(v).__name__, "len": len(v),
                    "items": [_describe(x, _depth + 1) for x in v]}
        return {"type": type(v).__name__, "repr": repr(v)}
    except Exception as e:
        return {"type": "DESCRIBE_ERROR", "error": str(e)}


def _dlog(event: str, payload: dict = None):
    """全イベントをJSONLに記録。seq/elapsed/呼び出し元を自動付与"""
    global _SEQ
    _SEQ += 1
    rec = {"seq": _SEQ, "ts": datetime.datetime.now().isoformat(),
           "elapsed": round(time.time() - _START_TIME, 4), "event": event}
    try:
        import inspect
        fr = inspect.stack()[1]
        rec["src"] = f"{fr.function}:{fr.lineno}"
    except Exception:
        pass
    rec.update(payload or {})
    os.makedirs("tmp", exist_ok=True)
    with open(_DEBUG_PATH, "a", encoding="utf-8") as f:
        f.write(json.dumps(rec, ensure_ascii=False, default=str) + "\n")


def _dlog_full(event: str, context: str, value, **extra):
    """任意オブジェクトを完全記述して記録する"""
    _dlog(event, {"context": context, "full": _describe(value), **extra})


# ============================================================
# llm_call ラッパー（v30の _llm_call_safe と同方式）
# ============================================================


def _rescue_json(text: str, context: str = ""):
    """textから複数戦略でJSONを救出。各戦略の試行と結果をログに残す"""
    text = (text or "").strip()
    strategies = [("raw_loads", text)]
    m = re.search(r"```(?:json)?\s*(.*?)```", text, re.DOTALL)
    if m:
        strategies.append(("strip_fence", m.group(1).strip()))
    s, e = text.find("{"), text.rfind("}")
    if s != -1 and e > s:
        strategies.append(("brace_greedy", text[s:e + 1]))
    for name, cand in strategies:
        try:
            parsed = json.loads(cand)
            _dlog("rescue_attempt", {"context": context, "strategy": name, "ok": True})
            return parsed
        except Exception as pe:
            _dlog("rescue_attempt", {"context": context, "strategy": name,
                                     "ok": False, "error": str(pe)})
    _dlog("rescue_all_failed", {"context": context, "text": text})
    return None


async def _llm_call_safe(prompt: str, schema: dict = None, context: str = "",
                         pdf_path: str = "", model: str = "") -> dict:
    """llm_callのラッパー。入力・生レスポンス・fallback・救出の全段を記録する

    pdf_path を渡すと file_paths でPDF書類のまま Gemini に読ませる（モデル固定）。
    schema指定時は data を返す。失敗時はschemaなしfallback→JSON救出まで試みる。
    """
    from agent_sdk import llm_call

    t0 = time.time()
    cleaned = None
    if schema:
        def _clean(s):
            if isinstance(s, dict):
                s.pop("additionalProperties", None)
                for v in s.values():
                    _clean(v)
            elif isinstance(s, list):
                for v in s:
                    _clean(v)
        cleaned = copy.deepcopy(schema)
        _clean(cleaned)

    _dlog("llm_call_input_full", {"context": context, "prompt": prompt,
          "has_schema": schema is not None, "pdf_path": pdf_path, "model": model})

    kw = {"prompt": prompt}
    if cleaned:
        kw["schema"] = cleaned
    if pdf_path:
        kw["file_paths"] = [pdf_path]
        kw["model"] = GEMINI_PDF_MODEL
    elif model:
        kw["model"] = model

    res = await llm_call(**kw)
    _dlog_full("llm_call_raw_response", context, res, attempt=1)

    # 添付サイズ超過は決定論的エラー: 同じPDFの再試行は無意味なので即raise
    if pdf_path and isinstance(res, dict) and res.get("error"):
        err_s = str(res.get("error", ""))
        if ("inline limit" in err_s or "over the inline" in err_s
                or ("attachments total" in err_s and "bytes" in err_s)):
            raise PdfTooLargeError(err_s)

    # エラーならschemaなしでfallback（JSONのみ返させて救出）
    if isinstance(res, dict) and res.get("error") and schema:
        _dlog("llm_call_fallback_trigger", {"context": context,
              "error_full": str(res.get("error", ""))})
        fb_kw = {"prompt": prompt + "\n\n出力は有効なJSONオブジェクトだけ。"
                                    "説明・markdown・前置きは一切禁止。"}
        if pdf_path:
            fb_kw["file_paths"] = [pdf_path]
            fb_kw["model"] = GEMINI_PDF_MODEL
        elif model:
            fb_kw["model"] = model
        res = await llm_call(**fb_kw)
        _dlog_full("llm_call_raw_response", context, res, attempt=2)

    _dlog("llm_call_finish", {"context": context, "duration": round(time.time() - t0, 3)})

    if schema:
        data = res.get("data") if isinstance(res, dict) else None
        if data is None and isinstance(res, dict) and "text" in res:
            data = _rescue_json(res.get("text") or "", context)
        # 構造化出力の失敗は {"data": {"error": ...}} で返る（トップにerrorが出ない）
        if isinstance(data, dict) and "error" in data and len(data) <= 2:
            _dlog("llm_call_structured_error", {"context": context, "data": data})
            data = None
        if data is None:
            raise Exception(f"LLMの構造化出力が取得できませんでした。context={context}")
        _dlog_full("llm_call_return", context, data)
        return data
    return {"text": res.get("text", "") if isinstance(res, dict) else str(res),
            "error": res.get("error") if isinstance(res, dict) else None}


# ============================================================
# RAGツールのリトライラッパー（v30の call_with_retry と同方式・簡易版）
# ============================================================


async def _rag_call_with_retry(func, *args, max_attempts: int = 3,
                               base_wait: float = 1.0, retry_context: str = "",
                               **kwargs) -> dict:
    """rag_file_list / rag_download のリトライ。全試行を記録する"""
    last_exc = None
    for attempt in range(1, max_attempts + 1):
        t0 = time.time()
        try:
            res = await func(*args, **kwargs)
            _dlog_full("rag_raw_response", f"{retry_context}#attempt{attempt}", res,
                       dur=round(time.time() - t0, 3))
            if isinstance(res, dict) and res.get("error"):
                raise RuntimeError(f"tool error: {res['error']}")
            return res
        except Exception as e:
            last_exc = e
            _dlog("rag_attempt_failed", {"rc": retry_context, "attempt": attempt,
                  "error_type": type(e).__name__, "error_full": str(e)})
            if attempt == max_attempts:
                raise
            await asyncio.sleep(base_wait * (2 ** (attempt - 1)))
    raise last_exc  # 到達しない（保険）


# ============================================================
# ユーティリティ
# ============================================================


def _nfkc(s: str) -> str:
    """照合用の正規化（全角半角・大小文字・空白を吸収）"""
    return unicodedata.normalize("NFKC", str(s)).lower().replace(" ", "")


def _match_key(s: str) -> str:
    """ファイル名照合用の決定的キー（完全一致用）

    AgentPlatform上のファイル名は濁点・半濁点が分解形（結合文字U+3099/U+309Aや
    独立記号）で保存されることがある（一覧APIの実データで確認）。
    両側に同じ変換をかけてから完全一致で照合する。曖昧マッチはしない。
    """
    s = str(s).replace("゛", "゙").replace("゜", "゚")  # ゛゜→結合文字
    s = s.replace("ﾞ", "゙").replace("ﾟ", "゚")        # ﾞﾟ（半角）→結合文字
    return unicodedata.normalize("NFKC", s).replace(" ", "").lower()


def _ensure_openpyxl() -> None:
    """openpyxl が無い環境ではサブプロセスの pip で導入する

    プロセス内で pip._internal を呼ぶ方式は禁止（以降の import に警告が連鎖し
    stdout 80KB を食い潰す実害があるため。ZETAメディアDBで観測済み）。
    """
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        import subprocess
        import sys
        subprocess.run([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl"],
                       capture_output=True, timeout=120)
        import openpyxl  # noqa: F401


# ============================================================
# 名前→UUID解決（rag_file_list 全件一覧との完全一致照合。検索は使わない）
# ============================================================


async def _load_listing() -> dict[str, str]:
    """rag_file_list をページングして全ファイルの {照合キー: uuid} を作る

    rag_search（中身検索）は使わない（ユーザー決定）。中身が検索インデックスに
    載らないファイル（xlsx・文字化けPDF・画像PDF）も一覧には必ず載るため、
    索引Excelに書かれた名前との完全一致照合だけで全ファイルを解決できる。
    """
    global _listing_cache
    if _listing_cache is not None:
        return _listing_cache
    from agent_sdk import rag_file_list

    listing: dict[str, str] = {}
    offset = 0
    while offset <= 1000:  # 暴走ガード（実データは80件弱）
        # 空リスト指定＝ルームに接続された学習データ全部が対象（このルームは1つだけ）
        res = await _rag_call_with_retry(rag_file_list, [],
                                         limit=_LIST_PAGE, offset=offset,
                                         retry_context=f"file_list:offset{offset}")
        files = [f for g in res.get("results", []) or []
                 for f in g.get("files", []) or []]
        for f in files:
            if f.get("name") and f.get("uuid"):
                listing[_match_key(f["name"])] = f["uuid"]
        if len(files) < _LIST_PAGE:
            break
        offset += _LIST_PAGE
    if not listing:
        raise Exception(
            "rag_file_list が0件を返しました（INDEX_UUID と学習データの接続を確認）")
    _dlog("listing_loaded", {"count": len(listing)})
    _listing_cache = listing
    return listing


async def _collect_uuids_by_names(names: list[str]) -> dict[str, str]:
    """ファイル名→uuid を解決する（全件一覧との完全一致照合）

    索引Excel自身も一覧の名前で解決するため、Excelを差し替えてUUIDが変わっても
    自動で追従する。一覧で引けない場合のみ INDEX_XLSX_UUID にフォールバックする。
    """
    listing = await _load_listing()
    for n in names:
        key = _match_key(n)
        if key in _uuid_cache:
            continue
        if key in listing:
            _uuid_cache[key] = listing[key]
        elif key == _match_key(INDEX_XLSX_NAME):
            _uuid_cache[key] = INDEX_XLSX_UUID  # 一覧に索引Excelが載らない環境向けの保険
    return dict(_uuid_cache)


async def _download_files(names: list[str], *, strict: bool = True) -> dict[str, str]:
    """指定名のファイルを rag_download し {元の名前: sandbox_path} を返す

    名前解決できないファイルは strict=True なら説明付き例外、strict=False なら
    結果から除外して続行する（呼び出し側が資料単位のエラーとして扱う）。
    """
    from agent_sdk import rag_download

    result: dict[str, str] = {}
    need_dl: list[tuple[str, str]] = []  # (name, uuid)
    uuids = await _collect_uuids_by_names(names)
    unresolved: list[str] = []
    for n in names:
        key = _match_key(n)
        if key in _path_cache and os.path.exists(_path_cache[key]):
            result[n] = _path_cache[key]
        elif key in uuids:
            need_dl.append((n, uuids[key]))
        else:
            unresolved.append(n)
    if unresolved:
        if strict:
            raise Exception(
                f"学習データに見つからないファイルがあります: {unresolved}。"
                "索引ExcelのPDF名と学習データの実ファイル名が一致しているか確認してください。")
        _dlog("unresolved_skipped", {"names": unresolved})
        print(f"  WARN: 一覧に見つからないファイルをスキップ: {unresolved}")
    if need_dl:
        dl = await _rag_call_with_retry(rag_download, [u for _, u in need_dl],
                                        retry_context="download")
        downloaded = {d["uuid"]: d["sandbox_path"] for d in dl.get("downloaded_files", [])}
        for n, u in need_dl:
            if u not in downloaded:
                raise Exception(f"{n} のダウンロードに失敗しました: {dl.get('errors')}")
            _path_cache[_match_key(n)] = downloaded[u]
            result[n] = downloaded[u]
    _dlog("download_files_done", {"resolved": list(result), "unresolved": unresolved})
    return result


# ============================================================
# 内部: 索引Excel読解
# ============================================================


async def _load_index() -> list[dict]:
    """索引Excel（商材DBシート）を読み、全商材の行を list[dict] で返す（キャッシュあり）"""
    global _index_cache
    if _index_cache is not None:
        return _index_cache
    _ensure_openpyxl()
    paths = await _download_files([INDEX_XLSX_NAME])
    import openpyxl

    wb = openpyxl.load_workbook(paths[INDEX_XLSX_NAME], data_only=True)
    ws = wb["商材DB"]
    rows: list[dict] = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name:
            continue
        rows.append({
            "商材名": str(name),
            "商材カテゴリ": str(ws.cell(r, 2).value or ""),
            "対象業種": str(ws.cell(r, 3).value or ""),
            "商材概要": str(ws.cell(r, 4).value or ""),
            "関連PDF": [p for p in str(ws.cell(r, 5).value or "").split(";") if p],
        })
    if not rows:
        raise Exception("索引Excelから商材行を1件も読めませんでした（シート名・列構成を確認）")
    _dlog("index_loaded", {"rows": len(rows)})
    _index_cache = rows
    return rows


def _find_row(rows: list[dict], product: str) -> dict | None:
    """商材名を索引の行に突き当てる（完全一致→包含の順）"""
    key = _nfkc(product)
    if not key:
        return None
    for r in rows:
        if _nfkc(r["商材名"]) == key:
            return r
    for r in rows:
        rk = _nfkc(r["商材名"])
        if key in rk or rk in key:
            return r
    return None


def _doc_kind(fname: str) -> str:
    """ファイル名から資料の種別を判定する

    customer=お客様に見せられる資料 / ops=操作・研修資料 /
    internal=社内資料（概要書・営業マニュアル。NET価格入り） / other=判別不能
    """
    f = _nfkc(fname)
    if "営業マニュアル" in f or "eigyomanual" in f:
        return "internal"
    if any(k in f for k in ("usermanual", "trainingbook", "操作マニュアル", "研修")):
        return "ops"
    if "概要書" in f or "概要資料" in f:
        return "internal"
    if any(k in f for k in ("提案書", "teiansho", "proposal", "チラシ", "flyer", "ポスター")):
        return "customer"
    return "other"


def _pick_files(row: dict, purpose: str, *, max_files: int = MAX_READ_FILES) -> list[str]:
    """商材と資料の用途から、実際に読むPDFをコードで決定する

    関連PDFは新しい版が先頭に並んでいる規約なので、リスト順＝新しい順で採用し、
    年月プレフィックスを除いて同一の資料（新旧版）は新しい方だけを使う。
    「顧客提示」では customer 以外の種別にフォールバックしない
    （社内向け資料をお客様向けとして選ばないことをコードで保証）。
    """
    if purpose == "顧客提示":
        kind_order = ["customer"]
    elif purpose == "操作・設定":
        kind_order = ["ops", "internal"]
    else:  # 仕様・価格・その他（不明な値もここに倒す）
        kind_order = ["internal", "customer", "ops", "other"]

    picked: list[str] = []
    seen_base: set[str] = set()
    for kind in kind_order:
        for name in row["関連PDF"]:
            if len(picked) >= max_files:
                return picked
            if _doc_kind(name) != kind or name in picked:
                continue
            base = _nfkc(re.sub(r"^\d{6}_", "", name))  # 年月を除いた同一資料判定
            if base in seen_base:
                continue
            seen_base.add(base)
            picked.append(name)
    return picked


# ============================================================
# 公開関数（メインAIが処理フローに沿って呼ぶ）
# ============================================================


async def get_categories() -> dict[str, int]:
    """① 商材カテゴリの一覧を返す {カテゴリ名: 商材数}"""
    rows = await _load_index()
    cats: dict[str, int] = {}
    for r in rows:
        cats[r["商材カテゴリ"]] = cats.get(r["商材カテゴリ"], 0) + 1
    _dlog("get_categories", {"categories": cats})
    return cats


async def get_products(category: str = None, industry: str = None) -> list[dict]:
    """② 商材の一覧を返す（カテゴリ・対象業種で絞り込み可。Noneなら全件）

    戻り値の各要素: {"商材名", "商材カテゴリ", "対象業種", "商材概要"}
    ※関連PDFのファイル名は返さない（PDFの決定は get_pdfs＝コードの責務）
    """
    rows = await _load_index()
    hits = rows
    if category:
        c = _nfkc(category)
        hits = [r for r in hits if c in _nfkc(r["商材カテゴリ"])]
    if industry:
        i = _nfkc(industry)
        hits = [r for r in hits if i in _nfkc(r["対象業種"])]
    _dlog("get_products", {"category": category, "industry": industry,
                           "hits": [r["商材名"] for r in hits]})
    return [{k: r[k] for k in ("商材名", "商材カテゴリ", "対象業種", "商材概要")}
            for r in hits]


async def get_pdfs(product: str, purpose: str = "仕様・価格・その他") -> list[str]:
    """③ 商材名と用途から、読むべきPDF名をコードが決定して返す

    Args:
        product: 商材名（get_products の「商材名」を一字一句そのまま渡す）
        purpose: 顧客提示 / 操作・設定 / 仕様・価格・その他
            ・顧客提示 → 提案書・チラシのみ返す（社内向け資料は絶対に返さない）
            ・操作・設定 → 操作マニュアル・トレーニングブック優先
            ・仕様・価格・その他 → 概要書（最も詳しい）優先

    Returns:
        PDF名のリスト（最大3件・新しい版優先）。商材が見つからなければ例外。
        顧客提示で客用資料が無い場合は空リスト（正直に「無い」と答えること）。
    """
    rows = await _load_index()
    row = _find_row(rows, product)
    if row is None:
        raise Exception(f"商材「{product}」が索引に見つかりません。"
                        "get_products の商材名を一字一句そのまま渡してください。")
    files = _pick_files(row, purpose)
    _dlog("get_pdfs", {"product": row["商材名"], "purpose": purpose, "files": files})
    return files


async def read_docs(question: str, pdf_names: list[str]) -> list[dict]:
    """④ 指定PDFを取り寄せ、1件ずつ llm_call（Gemini添付）で全文読解して回答を返す

    複数PDFは並列で読む。1件の失敗で全体を止めず、失敗分は {"file", "error"} で返す。
    Returns: [{"file": PDF名, "answer": 回答テキスト} または {"file", "error"}]
    """
    paths = await _download_files(pdf_names, strict=False)

    async def _read_one(name: str, path: str) -> dict:
        size = os.path.getsize(path)
        if size > MAX_PDF_BYTES:
            return {"file": name, "error": f"サイズ超過（{size // (1024 * 1024)}MB）"}
        prompt = f"""あなたはEPSILONの商材資料を読んで営業担当の質問に答えるアシスタントです。
添付PDF「{name}」だけを根拠に、次の質問に日本語で答えてください。

- この資料に記載がない事項は「この資料には記載なし」と明記する（推測・外部知識での補完は禁止）
- 価格に触れる場合は、定価とNET（社内仕切り値）の区別を資料の表記どおり明記する
- 回答の根拠となったページ番号を添える

## 質問
{question}
"""
        try:
            res = await _llm_call_safe(prompt, pdf_path=path, context=f"read:{name}")
        except PdfTooLargeError as e:
            return {"file": name, "error": f"添付サイズ上限超過: {e}"}
        except Exception as e:
            return {"file": name, "error": f"llm_call失敗: {e}"}
        if res.get("error"):
            return {"file": name, "error": str(res["error"])}
        return {"file": name, "answer": res.get("text", "")}

    readable = [n for n in pdf_names if n in paths]
    results = dict(zip(readable, await asyncio.gather(
        *(_read_one(n, paths[n]) for n in readable))))
    # 元の指定順を保ち、解決できなかった資料はerrorとして返す（隠さない）
    return [results.get(n, {"file": n, "error": "学習データの一覧に見つかりませんでした"
                                                "（索引Excelの名前と実ファイル名の不一致、または未登録）"})
            for n in pdf_names]
