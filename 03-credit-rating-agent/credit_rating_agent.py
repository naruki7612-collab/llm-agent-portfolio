from __future__ import annotations

import asyncio
import json
import logging
import os
import re
import shutil
from datetime import datetime, timezone, timedelta

import openpyxl

logger = logging.getLogger(__name__)

JST = timezone(timedelta(hours=9))
KONNEN = datetime.now(JST).year

LLM_MODEL = 'anthropic/claude-haiku-4-5'
ZENBUN_MODEL = LLM_MODEL


MATRIX_TABLE: dict[int, dict[str, str]] = {
    1: {'A': '正常', 'B': '要管理', 'C': '要管理', 'D': '正常（要注意）', 'E': '破綻懸念（要管理）',
        'F': '実質破綻（破綻懸念）', 'G': '実質破綻', 'H': '実質破綻', 'I': '破綻'},
    2: {'A': '要注意（正常2）', 'B': '要管理', 'C': '要管理', 'D': '要注意（要管理）',
        'E': '破綻懸念（要管理）', 'F': '実質破綻（破綻懸念）', 'G': '実質破綻', 'H': '実質破綻', 'I': '破綻'},
    3: {'A': '要注意', 'B': '要管理', 'C': '破綻懸念', 'D': '要注意（要管理）',
        'E': '破綻懸念（要管理）', 'F': '実質破綻（破綻懸念）', 'G': '実質破綻', 'H': '実質破綻', 'I': '破綻'},
    4: {'A': '要管理', 'B': '破綻懸念', 'C': '破綻懸念', 'D': '要管理（破綻懸念）',
        'E': '破綻懸念（要管理）', 'F': '実質破綻（破綻懸念）', 'G': '実質破綻', 'H': '実質破綻', 'I': '破綻'},
    5: {'A': '正常2', 'B': '要管理', 'C': '要管理', 'D': '要注意（要管理）',
        'E': '破綻懸念（要管理）', 'F': '実質破綻（破綻懸念）', 'G': '実質破綻', 'H': '実質破綻', 'I': '破綻'},
    6: {'A': '要注意', 'B': '要管理', 'C': '要管理', 'D': '要注意（要管理）',
        'E': '破綻懸念（要管理）', 'F': '実質破綻（破綻懸念）', 'G': '実質破綻', 'H': '実質破綻', 'I': '破綻'},
    7: {'A': '要注意（正常2）', 'B': '要管理', 'C': '要管理', 'D': '要注意（要管理）',
        'E': '破綻懸念（要管理）', 'F': '実質破綻（破綻懸念）', 'G': '実質破綻', 'H': '実質破綻', 'I': '破綻'},
    8: {'A': '要注意', 'B': '要管理', 'C': '要管理', 'D': '要注意（要管理）',
        'E': '破綻懸念（要管理）', 'F': '実質破綻（破綻懸念）', 'G': '実質破綻', 'H': '実質破綻', 'I': '破綻'},
    9: {'A': '正常（みなし）', 'B': '－', 'C': '－', 'D': '要注意（要管理）',
        'E': '破綻懸念（要管理）', 'F': '実質破綻（破綻懸念）', 'G': '実質破綻', 'H': '実質破綻', 'I': '破綻'},
    10: {'A': '正常（みなし）', 'B': '要管理（破綻懸念）', 'C': '要管理（破綻懸念）', 'D': '要注意（要管理）',
         'E': '破綻懸念', 'F': '実質破綻', 'G': '実質破綻', 'H': '実質破綻', 'I': '破綻'},
    11: {'A': '要管理（要注意）', 'B': '要管理', 'C': '破綻懸念（要管理）', 'D': '破綻懸念（要管理）',
         'E': '破綻懸念（要管理）', 'F': '実質破綻（破綻懸念）', 'G': '実質破綻', 'H': '実質破綻', 'I': '破綻'},
}

ROW_NAMES: dict[int, str] = {
    1: '債務超過なし・黒字かつ繰損なし', 2: '債務超過なし・赤字または繰損あり（注②）',
    3: '債務超過・黒字（注③）', 4: '債務超過・赤字',
    5: '実質債務償還年数10年超15年以内', 6: '実質債務償還年数15年超',
    7: '財務諸表のない先／企業集団で重要な関連会社の決算書のない先（注⑤）',
    8: '業歴2年未満の先', 9: '保証のみ先（注⑥）', 10: '不動産担保ローンのみ先（注⑦）',
    11: '外部情報懸念先（注⑧）',
}

COLUMN_NAMES: dict[str, str] = {
    'A': '延滞なし', 'B': '延滞1ヶ月以内・延滞懸念（注⑨）', 'C': '延滞2ヶ月以内', 'D': '条件変更（注⑫）',
    'E': '延滞3ヶ月以上6ヶ月未満・倒産懸念（注⑩）', 'F': '延滞6ヶ月以上12ヶ月未満（注⑪）',
    'G': '延滞12ヶ月以上', 'H': '個別貸倒引当等', 'I': '法的整理・取引停止処分',
}

CHUKI_ZENBUN = """\
注① マトリックス該当分を低位区分にて判定
注② 赤字とは、営業利益・経常利益・当期利益のいずれかが赤字の場合。
    赤字でも一過性・キャッシュフロー・改善計画等により正常先2とする場合あり
注③ 債務超過の解消に10年以上要する債務者は要管理とする場合あり。
    代表者等からの借入金等は自己資本にみなすことは可
注④ 実質債務償還年数＝（短期借入金＋長期借入金＋その他有利子負債（リース割賦他）
    −劣後債務−正常運転資金）÷（税引後利益＋減価償却費）
    正常運転資金＝受取手形＋売掛金＋棚卸資産−流動性支払手形−買掛金
注⑤ 個人事業者で貸借対照表のない場合は「番号7」にて判定
注⑥ 提携保証のみ先は保証会社の債務者区分に合わせる（みなし正常先）。
    但し外部懸念情報等で「要注意先」以下とする場合あり
注⑦ 不動産担保ローン（定型商品・保全確保）先は財務内容・損益状況にかかわらず
    正常先（みなし）とする
注⑧ 外部情報懸念先（債務者の営業実態により「要管理先」とせず「要注意先」とする場合もある）
    ・過去5年以内にブラック情報歴のある先
    ・過去5年以内に不動産差押、仮差押歴のある先
    ・取引先に大口焦付（年商×10％以上）が発生した先
    ・取引先に金利減免または返済条件の緩和の申し入れをしている、
      または当社への支払いにおいて危惧される事情が発生した先
    ・取引先とのトラブル・外部情報・融手操作・高利借入等
注⑨ 延滞懸念〜2か月連続自振不能・月越延滞（1か月以内）等、支払振りに問題のある先
注⑩ 倒産懸念〜管理債権計上基準先
    （債務者の営業実態により「破綻懸念先」とせず「要管理先」とする場合もある）
注⑪ 6か月以上延滞先でも債務者の営業実態により「実質破綻先」とせず
    「破綻懸念先」とする場合もある
注⑫ 当初期日から2年以上経過した先、および3回以上条件変更を実施した先は
    要管理以下先とする場合もある。但し直近条件変更が当初条件と同一以上に改善の場合、
    変更後の返済条件が妥当かつ履行が確実な場合は除く"""

TOKUREI_NO_BANGO = '②③⑥⑦⑧⑩⑪⑫'
NON_TOKUREI_SETSUMEI = (
    '注①＝複数の縦軸に該当した場合に低位区分を採るという「判定ルール」、'
    '注④＝実質債務償還年数の「算式の定義」、'
    '注⑤＝個人事業者の行選択ルール、'
    '注⑨＝横軸B（延滞懸念）の定義。'
    'これら4つは特例ではないため、「特例適用」欄に書いてはならない。'
)

KUBUN_ORDER = ['正常', '正常2', '要注意', '要管理', '破綻懸念', '実質破綻', '破綻']

CRITICAL_BS_ITEMS = ['総資産', '自己資本', '有利子負債合計']
CRITICAL_PL_ITEMS = ['営業利益', '経常利益', '税引後利益', '減価償却費']


KARTE_SHEET_GYOKYO = '業況推移表'
KARTE_SHEET_KIGYO = '企業概要表'
KARTE_SHEET_HYOSHI = '顧客カルテ'
KARTE_SHEET_GROUP = 'グループ合併バランス'

ITEM_CODE_COL = 1
PERIOD_HEADER_ROW = 5
PERIOD_MIN_COL = 10
TOKKI_COL = 20

ZAIMU_ITEM_CODES: dict[str, int] = {
    '流動資産': 10, '受取手形': 70, '売掛金': 80, '棚卸資産': 120,
    '総資産': 430,
    '流動負債': 440, '流動性支払手形': 450, '買掛金': 470,
    '短期借入金': 490,
    '一年内返済長期借入金': 540,
    'その他有利子負債_流動': 560,
    '長期借入金': 630,
    'その他有利子負債_固定': 640,
    '自己資本': 690, '資本金': 700, '繰越利益剰余金': 720,
    '有利子負債合計': 730,
    '売上高': 740, '月商': 750, '売上原価': 760, '売上総利益': 780,
    '販管費': 800, '営業利益': 820, '支払利息': 830, '経常利益': 840,
    '特別利益': 850, '特別損失': 860, '税引後利益': 870,
    '減価償却費': 890, '役員報酬': 910, '割引・譲渡手形': 960,
    '劣後債務': 980,
}

KARTE_SHIHYO_CODES: dict[str, int] = {
    '自己資本比率': 990, '流動比率': 1000, '売上債権回転期間': 1010,
    '買入債務回転期間': 1020, '正常運転資金': 1030, '実質債務': 1040,
    '実質債務償還年数': 1050, '金融費用率': 1060,
}

NON_KINGAKU_ITEMS = frozenset({
    '自己資本比率', '流動比率', '売上債権回転期間', '買入債務回転期間',
    '実質債務償還年数', '金融費用率',
})
PERCENT_ITEMS = frozenset({'自己資本比率', '流動比率', '金融費用率'})
SENYEN_TO_HYAKUMAN = 1000

KIGYO_ANCHOR_COL = 2
KIGYO_ANCHORS: dict[str, str] = {
    '商号・氏名': '商号', '業種': '業種', '沿革': '沿革',
    '経営者': '経営者', '業績': '業績', '事業所': '事業所',
}

HYOSHI_VALUE_SEARCH_WIDTH = 6
HYOSHI_LABELS = (
    '取引先', 'グループ', '主要銀行', '支社／支店', '住所', '電話番号',
    '従業員数', '資本金', '業種', '設立日', '取引開始', '債務者区分', '分類',
    '代表者', '年齢',
)
GROUP_TOTAL_COL = 14
GROUP_SELF_COL = 15
GROUP_LABELS = ('債務者区分判定日', '区分', '延滞金')

OUTPUT_CELLS: dict[str, str] = {
    '対象先名': 'C2', '顧客番号': 'G2', '判定基準日': 'C3', '直近決算期': 'G3',
    '部店名': 'C4', '担当者': 'G4', '会社概要要約': 'B6',
    '縦軸': 'E10', '横軸': 'E11', '交点': 'E12', '特例': 'E13', '最終区分': 'E14', '判定根拠': 'E15',
    '業界特異性所見': 'B45',
    '参照ナレッジ': 'D48', '参照カルテ': 'D49', '使用モデル生成日時': 'D50',
}
KEIEI_SHIHYO_ROWS = {'売上高': 21, '営業利益': 22, '経常利益': 23, '税引後利益': 24,
                     '減価償却費': 25, '有利子負債合計': 26}
ZAIMU_SHIHYO_ROWS = {'総資産': 28, '自己資本': 29, '資本金': 30, '流動資産': 31,
                     '流動負債': 32, '棚卸資産': 33, '売掛金': 34, '実質債務': 35}
KEIEI_BUNSEKI_ROWS = {'自己資本比率': 37, '流動比率': 38, '売上債権回転期間': 39,
                      '買入債務回転期間': 40, '正常運転資金': 41, '実質債務償還年数': 42, '金融費用率': 43}
GYOKYO_BUNSEKI_ROWS = {**KEIEI_SHIHYO_ROWS, **ZAIMU_SHIHYO_ROWS, **KEIEI_BUNSEKI_ROWS}

NUMFMT_KINGAKU = '#,##0'
NUMFMT_SHOSU = '#,##0.##'
SHOSU_ITEMS = frozenset({
    '自己資本比率', '流動比率', '金融費用率',
    '売上債権回転期間', '買入債務回転期間',
    '実質債務償還年数',
})


class FormatMismatchError(Exception):
    pass


def _validate_label(ws, row: int, col: int, expected: str) -> None:
    actual = ws.cell(row, col).value
    if actual != expected:
        raise FormatMismatchError(
            f'[{ws.title}]({row},{col}) の見出しが期待値と不一致。'
            f'期待="{expected}" 実際="{actual}"。フォーマット不一致のため判定を停止します。'
        )


def _resolve_template_path(template_path: str) -> str:
    import glob

    if os.path.exists(template_path):
        return template_path

    search_dir = os.path.dirname(template_path) or 'rag'
    candidates = glob.glob(os.path.join(search_dir, '*.xlsx'))
    matched = [c for c in candidates
               if any(k in os.path.basename(c) for k in ('様式', '分析シート', '債務者区分'))]

    if len(matched) == 1:
        logger.warning(
            f'指定された様式ファイルが見つかりません（{template_path}）。'
            f'名前が近い「{matched[0]}」を使用します。'
            f'rag_downloadがファイル名の記号を置換した可能性があります。'
        )
        return matched[0]

    listing = os.listdir(search_dir) if os.path.isdir(search_dir) else '(ディレクトリ自体が存在しない)'
    raise Exception(
        f'出力様式のExcelが見つかりません: {template_path}\n'
        f'  {search_dir} の中身: {listing}\n'
        f'  候補が複数/ゼロのため自動選択しませんでした（該当候補: {matched}）。\n'
        f'  対処：様式ExcelがRAGに登録されているか確認し、rag_downloadの戻り値'
        f'（downloaded_files[].sandbox_path）をそのまま template_path に渡してください。'
        f'パスをファイル名から組み立てないこと。'
    )


class _XlsCell:

    __slots__ = ('value',)

    def __init__(self, value: object) -> None:
        self.value = value


class _XlsSheet:

    def __init__(self, sheet, datemode: int) -> None:
        self._s = sheet
        self._datemode = datemode

    @property
    def title(self) -> str:
        return self._s.name

    @property
    def max_row(self) -> int:
        return self._s.nrows

    @property
    def max_column(self) -> int:
        return self._s.ncols

    def cell(self, row: int, column: int) -> _XlsCell:
        import xlrd

        r, c = row - 1, column - 1
        if r < 0 or c < 0 or r >= self._s.nrows or c >= self._s.ncols:
            return _XlsCell(None)
        ctype = self._s.cell_type(r, c)
        value = self._s.cell_value(r, c)
        if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            return _XlsCell(None)
        if ctype == xlrd.XL_CELL_ERROR:
            return _XlsCell(None)
        if ctype == xlrd.XL_CELL_DATE:
            try:
                return _XlsCell(xlrd.xldate_as_datetime(value, self._datemode))
            except Exception:
                return _XlsCell(value)
        if ctype == xlrd.XL_CELL_TEXT:
            return _XlsCell(value if value != '' else None)
        if ctype == xlrd.XL_CELL_NUMBER and isinstance(value, float) and value.is_integer():
            return _XlsCell(int(value))
        return _XlsCell(value)


class _XlsWorkbook:

    def __init__(self, path: str) -> None:
        import xlrd

        self._b = xlrd.open_workbook(path)
        self._cache: dict[str, _XlsSheet] = {}

    @property
    def sheetnames(self) -> list[str]:
        return list(self._b.sheet_names())

    def __getitem__(self, name: str) -> _XlsSheet:
        if name not in self._cache:
            self._cache[name] = _XlsSheet(self._b.sheet_by_name(name), self._b.datemode)
        return self._cache[name]

    @property
    def worksheets(self) -> list[_XlsSheet]:
        return [self[n] for n in self.sheetnames]


def _load_karte_workbook(path: str):
    ext = os.path.splitext(path)[1].lower()
    if ext == '.xls':
        try:
            import xlrd
        except ImportError as e:
            raise Exception(
                f'旧xls形式のカルテを読むには xlrd が必要ですが、環境に入っていません（{e}）。'
                f'カルテを .xlsx 形式で出力し直してご提供いただくか、'
                f'Googleスプレッドシートから「Microsoft Excel (.xlsx)」でダウンロードしてください。'
            ) from e
        logger.info(f'旧xls形式のためxlrdで読み込みます: {os.path.basename(path)}')
        return _XlsWorkbook(path)
    return openpyxl.load_workbook(path, data_only=True)


def _to_float(v: object) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.replace(',', '').replace('　', '').strip()
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _detect_period_columns(ws) -> list[tuple[str, int]]:
    import re

    found: list[tuple[str, int]] = []
    for c in range(PERIOD_MIN_COL, ws.max_column + 1):
        v = ws.cell(PERIOD_HEADER_ROW, c).value
        if isinstance(v, str) and re.fullmatch(r'\d{4}/\d{1,2}', v.strip()):
            found.append((v.strip(), c))
        elif isinstance(v, datetime):
            found.append((v.strftime('%Y/%m'), c))
    if len(found) < 3:
        raise FormatMismatchError(
            f'[{ws.title}] {PERIOD_HEADER_ROW}行目から期のヘッダ（例「2025/03」）を'
            f'3期分以上検出できませんでした（検出：{found}）。'
            f'様式が想定と異なるため判定を停止します。'
        )
    return found


def _read_gyokyo_suii(ws) -> dict:
    all_periods = _detect_period_columns(ws)
    recent = all_periods[-3:]
    period_keys = ('前々期', '前期', '当期')
    col_of = {k: c for k, (_, c) in zip(period_keys, recent)}
    logger.info(f'業況推移表の期を検出：{[p for p, _ in all_periods]} → 直近3期'
                f'{[p for p, _ in recent]}を使用')

    code_to_row: dict[int, int] = {}
    for r in range(1, ws.max_row + 1):
        code = _to_float(ws.cell(r, ITEM_CODE_COL).value)
        if code is not None and code == int(code):
            code_to_row[int(code)] = r

    def _pull(codes: dict[str, int]) -> tuple[dict, list[str]]:
        out: dict[str, dict[str, float | None]] = {}
        miss: list[str] = []
        for name, code in codes.items():
            row = code_to_row.get(code)
            if row is None:
                miss.append(name)
                continue
            vals: dict[str, float | None] = {}
            for key in period_keys:
                raw = _to_float(ws.cell(row, col_of[key]).value)
                if raw is None:
                    vals[key] = None
                elif name in NON_KINGAKU_ITEMS:
                    vals[key] = round(raw * 100, 2) if name in PERCENT_ITEMS else round(raw, 2)
                else:
                    vals[key] = round(raw / SENYEN_TO_HYAKUMAN, 1)
            out[name] = vals
        return out, miss

    zaimu, missing = _pull(ZAIMU_ITEM_CODES)
    karte_shihyo, shihyo_missing = _pull(KARTE_SHIHYO_CODES)
    if shihyo_missing:
        logger.info(f'カルテ側で算出済みの指標のうち取得できなかったもの: {shihyo_missing}')

    if 'その他有利子負債_流動' in zaimu and 'その他有利子負債_固定' in zaimu:
        merged: dict[str, float | None] = {}
        for key in period_keys:
            a = zaimu['その他有利子負債_流動'][key]
            b = zaimu['その他有利子負債_固定'][key]
            merged[key] = None if a is None or b is None else round(a + b, 1)
        zaimu['その他有利子負債（リース割賦他）'] = merged

    tokki = [str(ws.cell(r, TOKKI_COL).value).strip()
             for r in range(1, ws.max_row + 1)
             if ws.cell(r, TOKKI_COL).value not in (None, '')]

    return {
        'periods': [p for p, _ in recent],
        'all_periods': [p for p, _ in all_periods],
        'zaimu': zaimu,
        'karte_shihyo': karte_shihyo,
        'missing': missing,
        'tokki': tokki,
    }


def _find_anchor_row(ws, text: str, col: int) -> int | None:
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.replace(' ', '').replace('　', '').replace('\n', '') == text:
            return r
    return None


def _read_kigyo_gaiyo(ws) -> dict:
    g: dict[str, object] = {}
    anchors = {name: _find_anchor_row(ws, name, KIGYO_ANCHOR_COL)
               for name in KIGYO_ANCHORS}
    missing_anchors = [n for n, r in anchors.items() if r is None]
    if '商号・氏名' in missing_anchors or '業種' in missing_anchors:
        raise FormatMismatchError(
            f'[{ws.title}] 列{KIGYO_ANCHOR_COL}に見出し「商号・氏名」「業種」が見つかりません'
            f'（検出できた見出し：{[n for n, r in anchors.items() if r]}）。'
            f'様式が想定と異なるため判定を停止します。'
        )

    r = anchors['商号・氏名']
    g['商号カナ'] = ws.cell(r, 4).value
    g['商号'] = ws.cell(r + 1, 4).value
    for rr in range(r, min(r + 8, ws.max_row) + 1):
        lab = ws.cell(rr, 4).value
        if isinstance(lab, str) and '顧客番号' in lab:
            g['顧客番号'] = ws.cell(rr, 8).value
            break

    r = anchors['業種']
    g['業種'] = ws.cell(r, 4).value
    shihonkin = _to_float(ws.cell(r, 16).value)
    g['資本金'] = None if shihonkin is None else round(shihonkin / SENYEN_TO_HYAKUMAN, 1)

    if anchors.get('沿革'):
        r = anchors['沿革']
        yy, mm = _to_float(ws.cell(r, 6).value), _to_float(ws.cell(r, 10).value)
        if yy:
            g['設立年'] = f'{int(yy)}年{int(mm)}月設立' if mm else f'{int(yy)}年設立'
        g['沿革'] = [str(ws.cell(rr, 4).value).strip()
                     for rr in range(r + 1, min(r + 6, ws.max_row) + 1)
                     if ws.cell(rr, 4).value not in (None, '')]

    if anchors.get('経営者'):
        r = anchors['経営者']
        later = [row for row in anchors.values() if row and row > r]
        limit = min(later) - 1 if later else min(r + 8, ws.max_row)
        yakuin = []
        for rr in range(r + 1, limit + 1):
            yaku, name = ws.cell(rr, 4).value, ws.cell(rr, 14).value
            if yaku and name:
                yakuin.append(f'{yaku}：{name}')
                if '代表' in str(yaku) and '代表者' not in g:
                    g['代表者'] = name
                    g['代表者年令'] = ws.cell(rr, 21).value
        g['役員'] = yakuin

    if anchors.get('業績'):
        r = anchors['業績']
        gyoseki = []
        for rr in range(r + 1, min(r + 8, ws.max_row) + 1):
            ki = ws.cell(rr, 4).value
            if isinstance(ki, str) and '/' in ki:
                gyoseki.append({
                    '期': ki.strip(),
                    '売上高': _to_float(ws.cell(rr, 8).value),
                    '経常利益': _to_float(ws.cell(rr, 14).value),
                    '税引後利益': _to_float(ws.cell(rr, 19).value),
                })
        g['業績'] = gyoseki
        if gyoseki:
            g['直近売上規模'] = gyoseki[-1]['売上高']

    if anchors.get('事業所'):
        r = anchors['事業所']
        g['本店所在地'] = ws.cell(r, 9).value
        g['特記事項'] = ws.cell(r, 25).value

    for rr in range(1, ws.max_row + 1):
        lab = ws.cell(rr, KIGYO_ANCHOR_COL).value
        if not isinstance(lab, str):
            continue
        flat = lab.replace('\n', '')
        if '取扱品目' in flat:
            g['取扱品目'] = [str(ws.cell(x, y).value).strip()
                          for x in range(rr, min(rr + 4, ws.max_row) + 1)
                          for y in (5, 14)
                          if ws.cell(x, y).value not in (None, '')]
        elif '大株主' in flat:
            g['大株主'] = [str(ws.cell(x, 5).value).strip()
                        for x in range(rr, min(rr + 5, ws.max_row) + 1)
                        if ws.cell(x, 5).value not in (None, '')]
        elif '上場' == flat:
            g['上場'] = ws.cell(rr, 4).value
    for rr in range(1, ws.max_row + 1):
        for cc in range(20, min(ws.max_column, 34) + 1):
            v = ws.cell(rr, cc).value
            if isinstance(v, str) and '外部懸念情報' in v:
                g['外部懸念情報'] = ws.cell(rr, cc + 5).value
                break

    return g


def _read_labeled_values(ws, labels: tuple[str, ...], width: int) -> dict[str, object]:
    out: dict[str, object] = {}
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            key = v.replace(' ', '').replace('　', '').replace('\n', '').strip()
            if key in labels and key not in out:
                for cc in range(c + 1, min(c + width, ws.max_column) + 1):
                    val = ws.cell(r, cc).value
                    if val not in (None, ''):
                        out[key] = val
                        break
    return out


def _read_hyoshi(ws) -> dict:
    info = _read_labeled_values(ws, HYOSHI_LABELS, HYOSHI_VALUE_SEARCH_WIDTH)

    entai_total: float | None = None
    yoshin_moto: float | None = None
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip() == '延滞金':
                for rr in range(r + 1, min(r + 10, ws.max_row) + 1):
                    if str(ws.cell(rr, 3).value).strip() == '合計':
                        entai_total = _to_float(ws.cell(rr, c).value)
                        break
                break
        if entai_total is not None:
            break
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 3).value).strip() == '合計':
            yoshin_moto = _to_float(ws.cell(r, 15).value)
            break

    info['延滞金合計'] = entai_total
    info['与信元本残高合計'] = yoshin_moto
    return info


def _read_group_balance(ws) -> dict:
    out: dict[str, object] = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if not isinstance(v, str):
            continue
        key = v.replace(' ', '').replace('　', '').strip()
        if key in GROUP_LABELS and key not in out:
            out[key] = ws.cell(r, GROUP_SELF_COL).value
            out[f'{key}_グループ合計'] = ws.cell(r, GROUP_TOTAL_COL).value
    return out


def _parse_entai_months(entai_tsuki: object, entai_kin: object) -> float | None:
    import re

    if isinstance(entai_tsuki, (int, float)):
        return float(entai_tsuki)
    if isinstance(entai_tsuki, str):
        m = re.search(r'(\d+(?:\.\d+)?)', entai_tsuki)
        if m:
            return float(m.group(1))
    return None


def _kubun_rank(row: int, col: str) -> int:
    principle = MATRIX_TABLE[row][col].split('（')[0]
    return KUBUN_ORDER.index(principle) if principle in KUBUN_ORDER else -1


KOJIN_MISHOKAKU_MSG = '個人事業者（個人概要表のみ）の抽出ロジックは未実装です。法人先のみ対応しています。'


def _karte_sheet_check(wb) -> None:
    if KARTE_SHEET_GYOKYO not in wb.sheetnames:
        raise FormatMismatchError(
            f'{KARTE_SHEET_GYOKYO}シートが見つかりません（実際のシート：{wb.sheetnames}）。'
            f'判定を行わず停止します。'
        )
    if KARTE_SHEET_KIGYO not in wb.sheetnames:
        if '個人概要表' in wb.sheetnames:
            raise Exception(KOJIN_MISHOKAKU_MSG)
        raise FormatMismatchError(
            f'{KARTE_SHEET_KIGYO}シートが見つかりません（実際のシート：{wb.sheetnames}）。'
        )


def _karte_yuko_hantei(karte_path: str) -> tuple[bool, str | None, list[str]]:
    """機械的なシート構成チェックのみ行う（数値抽出・AI呼び出しはしない）。

    複数の候補ファイルから有効な顧客カルテを選ぶための軽量プローブ。
    戻り値: (有効か, 無効な理由（有効なら None）, 実際のシート名一覧)
    """
    try:
        wb = _load_karte_workbook(karte_path)
    except Exception as e:
        return False, f'{type(e).__name__}: {e}', []
    try:
        _karte_sheet_check(wb)
    except Exception as e:
        return False, str(e), wb.sheetnames
    return True, None, wb.sheetnames


_KARTE_SETSUMEI_SCHEMA = {
    'type': 'object',
    'properties': {
        '推測されるファイルの正体': {
            'type': 'string',
            'description': '例：操作マニュアル、複数社分の内部集計資料、'
                           '別の顧客の一括印刷ファイル、正しいカルテだが一部シート名が違う 等'},
        'ユーザーへの説明文': {
            'type': 'string',
            'description': '1〜2文の日本語。このファイルが何に見えるか、'
                           'なぜ顧客カルテとして使えないかを、審査担当者にわかる言葉で説明する'},
    },
    'required': ['推測されるファイルの正体', 'ユーザーへの説明文'],
}


async def _karte_shippai_setsumei(karte_path: str, sheet_names: list[str],
                                  kikai_riyu: str) -> str:
    """機械チェックで弾かれたファイルについて、中身を軽く覗いてAIに自然文で説明させる。

    AIが使えない・失敗した場合は機械的な理由をそのまま返す（読めなくなることを防ぐ）。
    """
    try:
        from agent_sdk import llm_call
    except ImportError:
        return kikai_riyu

    sample = ''
    try:
        wb = _load_karte_workbook(karte_path)
        if wb.sheetnames:
            ws = wb[wb.sheetnames[0]]
            gyo = []
            for r in range(1, min(ws.max_row, 15) + 1):
                cells = [str(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 8) + 1)
                         if ws.cell(r, c).value not in (None, '')]
                if cells:
                    gyo.append('  '.join(cells))
            sample = '\n'.join(gyo)
    except Exception:
        pass

    prompt = f"""債務者区分判定のために顧客カルテ（Excel）が添付されましたが、
想定しているシート構成と一致せず、機械的なチェックで弾かれました。

【機械チェックの結果】{kikai_riyu}
【実際のシート名】{sheet_names}
【先頭シートの中身（一部）】
{sample[:1500] or '（読み取れませんでした）'}

このファイルが何であるかを中身から推測し、顧客カルテとして使えない理由を
審査担当者にわかりやすく1〜2文の日本語で説明してください。
"""
    try:
        res = await llm_call(prompt=prompt, schema=_KARTE_SETSUMEI_SCHEMA, model=LLM_MODEL)
        setsumei = (res.get('data') or {}).get('ユーザーへの説明文')
        return setsumei or kikai_riyu
    except Exception as e:
        logger.warning(f'ファイル判別の説明生成に失敗しました（機械的な理由をそのまま使います）: '
                       f'{type(e).__name__}: {e}')
        return kikai_riyu


async def select_karte(karte_candidates: list[str]) -> dict:
    """複数の候補ファイルから、有効な顧客カルテを1つ選ぶ。

    uploads/ には過去のやり取りの古いファイルも残り続けるため、「どれを使うか」を
    機械的なインデックスや更新日時だけで決めない。候補ごとに実際の中身（シート構成）を
    検証し、有効なものだけを対象にする。有効なものが複数あれば最も新しいものを使う
    （同名ファイルの再アップロードもこれで解決する）。有効なものが1つも無ければ、
    最も新しい候補の中身をAIに読ませ、何のファイルに見えるかを説明する。

    戻り値:
      {'karte_path': str}                              … 有効なファイルが見つかった
      {'error': str, '機械的な理由': str, 'ファイル': str} … 見つからなかった
    """
    junjo = sorted(karte_candidates, key=os.path.getmtime, reverse=True)
    yuko = []
    saishin_shippai = None
    for path in junjo:
        ok, riyu, sheets = _karte_yuko_hantei(path)
        if ok:
            yuko.append(path)
        elif saishin_shippai is None:
            saishin_shippai = (path, riyu, sheets)

    if yuko:
        return {'karte_path': yuko[0]}
    if saishin_shippai is None:
        return {'error': '有効な候補がありません'}

    path, riyu, sheets = saishin_shippai
    if riyu == KOJIN_MISHOKAKU_MSG:
        # 個人事業者のカルテは原因が明確なため、AIに読ませて説明させる必要が無い
        return {'error': riyu, '機械的な理由': riyu, 'ファイル': os.path.basename(path)}
    setsumei = await _karte_shippai_setsumei(path, sheets, riyu)
    return {'error': setsumei, '機械的な理由': riyu, 'ファイル': os.path.basename(path)}


async def extract_and_calculate(karte_path: str) -> dict:
    if not os.path.exists(karte_path):
        raise Exception(f'カルテファイルが見つかりません: {karte_path}')

    wb = _load_karte_workbook(karte_path)
    _karte_sheet_check(wb)
    ws_gaiyo = wb[KARTE_SHEET_KIGYO]
    ws_gyokyo = wb[KARTE_SHEET_GYOKYO]

    suii = _read_gyokyo_suii(ws_gyokyo)
    zaimu = suii['zaimu']
    karte_shihyo = suii['karte_shihyo']
    missing_items = suii['missing']
    if missing_items:
        logger.info(f'業況推移表に項目コードが見つからなかった財務項目: {missing_items}')

    gaiyo: dict[str, object] = _read_kigyo_gaiyo(ws_gaiyo)
    gaiyo['決算内容・特殊要因説明'] = suii['tokki']

    hyoshi: dict[str, object] = {}
    if KARTE_SHEET_HYOSHI in wb.sheetnames:
        hyoshi = _read_hyoshi(wb[KARTE_SHEET_HYOSHI])
    else:
        logger.info(f'{KARTE_SHEET_HYOSHI}シートがないため、延滞金は'
                    f'{KARTE_SHEET_GROUP}側の記載のみで判断します。')
    group: dict[str, object] = {}
    if KARTE_SHEET_GROUP in wb.sheetnames:
        group = _read_group_balance(wb[KARTE_SHEET_GROUP])
    gaiyo['表紙情報'] = hyoshi
    gaiyo['グループ情報'] = group
    for key in ('従業員数', '主要銀行', '設立日', '取引開始', '住所'):
        if key in hyoshi and key not in gaiyo:
            gaiyo[key] = hyoshi[key]

    meta: dict[str, object] = {
        '直近決算期': suii['periods'][-1],
        '部店': None,
        '担当者': None,
    }
    for key, label in (('部店', '部店名'), ('担当者', '担当者')):
        for r in range(1, min(6, ws_gyokyo.max_row) + 1):
            for c in range(1, ws_gyokyo.max_column + 1):
                v = ws_gyokyo.cell(r, c).value
                if isinstance(v, str) and v.strip() == label:
                    meta[key] = ws_gyokyo.cell(r, c + 1).value
                    break
            if meta[key]:
                break

    cur = {k: v['当期'] for k, v in zaimu.items()}

    missing_critical = [m for m in missing_items if m in CRITICAL_BS_ITEMS + CRITICAL_PL_ITEMS]
    empty_critical = [
        m for m in CRITICAL_BS_ITEMS + CRITICAL_PL_ITEMS
        if m not in missing_critical and cur.get(m) is None
    ]
    missing_critical += empty_critical
    if not gaiyo.get('業種'):
        missing_critical.append('業種')
    if missing_critical:
        return {
            'status': 'insufficient',
            'missing': missing_critical,
            'message': '判定に必要な情報が不足しているため、判定を行わず終了します。憶測での補完は行いません。',
        }

    seijou_unten_shikin = karte_shihyo.get('正常運転資金', {}).get('当期')
    shokan_nensu = karte_shihyo.get('実質債務償還年数', {}).get('当期')
    jisshitsu_saimu = karte_shihyo.get('実質債務', {}).get('当期')

    calc_unten = (
        (cur.get('受取手形') or 0) + (cur.get('売掛金') or 0) + (cur.get('棚卸資産') or 0)
        - (cur.get('流動性支払手形') or 0) - (cur.get('買掛金') or 0)
    )
    calc_bunshi = (cur.get('有利子負債合計') or 0) - (cur.get('劣後債務') or 0) - calc_unten
    bunbo = (cur.get('税引後利益') or 0) + (cur.get('減価償却費') or 0)
    kensan: dict[str, object] = {
        '正常運転資金_自前計算': round(calc_unten, 1),
        '実質債務_自前計算': round(calc_bunshi, 1),
        '算式分母（税引後利益＋減価償却費）': round(bunbo, 1),
    }
    if seijou_unten_shikin is not None and abs(calc_unten - seijou_unten_shikin) > 1:
        kensan['正常運転資金の差異'] = (
            f'カルテ{seijou_unten_shikin} vs 自前計算{round(calc_unten, 1)}（百万円）'
        )
        logger.info(f'正常運転資金がカルテ値と自前計算で相違：{kensan["正常運転資金の差異"]}')

    shokan_riyu = ''
    if shokan_nensu == 0:
        if calc_bunshi <= 0:
            shokan_riyu = (f'算出不能（実質債務が{round(calc_bunshi, 1)}百万円とマイナス＝'
                           f'正常運転資金が有利子負債を上回る実質無借金の状態のため）')
        elif bunbo <= 0:
            shokan_riyu = (f'算出不能（税引後利益{cur.get("税引後利益")}＋'
                           f'減価償却費{cur.get("減価償却費")}＝{round(bunbo, 1)}で分母が0以下のため）')
        else:
            shokan_riyu = '算出不能（カルテに0と記載。算出できなかったものとして扱う）'
        shokan_nensu = None
        logger.info(f'実質債務償還年数：カルテ記載が0のため{shokan_riyu}。縦軸5・6の判定は行わない')
    elif shokan_nensu is None:
        shokan_riyu = '算出不能（カルテに算出値がないため）'
        logger.info('実質債務償還年数：カルテ側に算出値がないため縦軸5・6の判定は行わない')

    saimu_choka = (cur.get('自己資本') or 0) < 0
    akaji = any((cur.get(k) or 0) < 0 for k in CRITICAL_PL_ITEMS[:3])
    kurison = (cur.get('繰越利益剰余金') or 0) < 0
    akaji_matawa_kurison = akaji or kurison

    if not saimu_choka and not akaji_matawa_kurison:
        row_no = 1
    elif not saimu_choka:
        row_no = 2
    elif not akaji:
        row_no = 3
    else:
        row_no = 4

    if kurison and not akaji:
        logger.info(
            f'繰越利益剰余金{cur.get("繰越利益剰余金")}百万円（繰損）を検知。'
            f'当期は黒字だが別表1の縦軸1は「黒字かつ繰損なし」が条件のため縦軸{row_no}とする'
        )

    candidate_rows = [row_no]
    if shokan_nensu is not None:
        if 10 < shokan_nensu <= 15:
            candidate_rows.append(5)
        elif shokan_nensu > 15:
            candidate_rows.append(6)

    entai_hyoshi = _to_float(hyoshi.get('延滞金合計'))
    entai_group = _to_float(group.get('延滞金'))
    entai_values = [v for v in (entai_hyoshi, entai_group) if v is not None]
    entai_kin = max(entai_values) if entai_values else None
    entai_tsuki = None
    yoko_note = ''

    if not entai_values:
        yoko_kigo = None
        yoko_note = ('延滞金の記載を顧客カルテ・グループ合併バランスのいずれからも読み取れませんでした。'
                     '延滞なしとみなすと区分を誤る可能性があるため、判定を行いません。')
        logger.info(yoko_note)
    elif entai_kin == 0:
        yoko_kigo = 'A'
        yoko_note = (f'延滞金0円（顧客カルテ={entai_hyoshi} / グループ合併バランス={entai_group}／単位：円）'
                     f'のため延滞なし（横軸A）と判定。'
                     f'なお条件変更履歴の欄がカルテにないため、横軸D（条件変更・注⑫）の'
                     f'該当有無は確認できていない。')
    else:
        yoko_kigo = None
        yoko_note = (f'延滞金{entai_kin:,.0f}円を検知しましたが、カルテに延滞月数の欄がないため'
                     f'横軸（B〜G）を特定できません。別表1の横軸は延滞月数で分岐するため、'
                     f'金額からの逆算は行いません。担当者への確認が必要です。')
        logger.info(yoko_note)

    if yoko_kigo is None:
        return {
            'status': 'insufficient',
            'missing': ['延滞月数（横軸A〜Iの判定に必要。実カルテに該当欄がありません）'],
            'message': yoko_note,
            'gaiyo': gaiyo, 'meta': meta, 'zaimu': zaimu,
        }

    best_row = max(candidate_rows, key=lambda r: _kubun_rank(r, yoko_kigo))
    kouten = MATRIX_TABLE[best_row][yoko_kigo]

    return {
        'status': 'ok',
        'gaiyo': gaiyo,
        'meta': meta,
        'zaimu': zaimu,
        'karte_shihyo': karte_shihyo,
        'missing_items': missing_items,
        'keisan': {
            '正常運転資金': seijou_unten_shikin,
            '実質債務': jisshitsu_saimu,
            '実質債務償還年数': shokan_nensu,
            '実質債務償還年数_算出不能理由': shokan_riyu,
            '算式分母': bunbo,
            '検算': kensan,
        },
        'hantei': {
            '該当行候補': candidate_rows,
            '採用行': best_row,
            '採用行名': ROW_NAMES[best_row],
            '横軸記号': yoko_kigo,
            '横軸名': COLUMN_NAMES[yoko_kigo],
            '横軸備考': yoko_note,
            '交点区分': kouten,
            '低位区分優先を適用したか': len(candidate_rows) > 1,
            '縦軸判定内訳': {
                '債務超過': saimu_choka,
                '赤字': akaji,
                '繰損': kurison,
            },
        },
    }


GYOKAI_KOTOWARI = ('※以下の【一般論】は業界一般の傾向です。出典・年次の妥当性は未検証のため、'
                   'ご確認のうえご判断ください。')


SHIRYO_SETSUMEI_ATAMA = ('本資料は', '本調査は', '当調査は', '当資料は', 'この資料は', 'この調査は')


def _shoken_kumitate(karte_items: list, gaibu_items: list,
                     zenbun: list) -> tuple[str, set[int], list[dict], dict]:
    """カルテ由来（別ターン）と外部情報由来（別ターン）の所見要素を、
    タグ付きの文章に組み立てる。

    どちらのターンから来たかで出典が決まる（AIに申告させない）。
    2つのターンを分けているため、カルテのターンは外部資料を一切見ていない。
    「外部情報がカルテ由来として紛れ込む」ことが構造的に起きないため、
    以前あった数値一致による二重チェックはここでは不要（別ターンでの残置チェックとして
    出典資料番号の実在チェックだけ残す）。

    カルテ側にも外部情報側にも【推測】【意見】が出るため、タグだけでは
    どちら由来か紛らわしい。見出し（「カルテ」「外部情報」）で区切って区別する。
    どちらか一方しか無い場合は見出しを付けない（区別する必要が無いため）。

    Returns:
        (組み立てた文章, 実際に引用された資料のno集合, 書き直しが必要な要素, 検査の記録)
    """
    ban = {z['no']: z for z in (zenbun or [])}
    karte_bun, gaibu_bun, mondai, gaibu_no = [], [], [], set()
    kaz = {'カルテ': 0, '外部情報': 0, 'カルテ_事実': 0, '外部情報_一般論': 0}

    for it in (karte_items or []):
        honbun = str(it.get('文') or '').strip().rstrip('。')
        if not honbun:
            continue
        shurui = it.get('種類') if it.get('種類') in ('事実', '推測', '意見') else '意見'
        full = f'{honbun}【{shurui}】。'
        karte_bun.append(full)
        kaz['カルテ'] += 1
        if shurui == '事実':
            kaz['カルテ_事実'] += 1

    for it in (gaibu_items or []):
        honbun = str(it.get('文') or '').strip().rstrip('。')
        if not honbun:
            continue
        shurui = it.get('種類') if it.get('種類') in ('一般論', '推測', '意見') else '一般論'
        full = f'{honbun}【{shurui}】。'
        gaibu_bun.append(full)
        kaz['外部情報'] += 1
        if shurui == '一般論':
            kaz['外部情報_一般論'] += 1
        no = it.get('出典資料番号')
        if no in ban:
            gaibu_no.add(no)
        else:
            mondai.append({'文': full, '理由': f'出典資料番号{no}が資料一覧に無い'})
        if honbun.startswith(SHIRYO_SETSUMEI_ATAMA):
            mondai.append({'文': full, '理由': '資料の説明にとどまり業界の傾向を述べていない'})

    if kaz['カルテ'] > 0 and kaz['カルテ_事実'] < 1:
        _glog('所見にカルテ由来の【事実】が1つも無い（枠の指示が守られていない可能性）')
    if kaz['外部情報'] > 0 and kaz['外部情報_一般論'] < 1:
        _glog('所見に外部情報由来の【一般論】が1つも無い（枠の指示が守られていない可能性）')

    if karte_bun and gaibu_bun:
        out = 'カルテ\n' + '\n'.join(karte_bun) + '\n\n外部情報\n' + '\n'.join(gaibu_bun)
    else:
        out = '\n'.join(karte_bun + gaibu_bun)

    kensa = {'カルテ由来の文数': kaz['カルテ'], '外部情報由来の文数': kaz['外部情報']}
    return out, gaibu_no, mondai, kensa


_SHUSEI_SCHEMA = {
    'type': 'object',
    'properties': {
        '修正': {
            'type': 'array',
            'description': '渡された文と同じ順番・同じ件数で返す',
            'items': {
                'type': 'object',
                'properties': {
                    '修正後の文': {'type': 'string', 'description':
                              '直した文。文末にタグを1つだけ付ける。'
                              '業界の傾向を述べられないなら「削除」の1語だけを返す'},
                },
                'required': ['修正後の文'],
            },
        },
    },
    'required': ['修正'],
}


async def _shoken_shusei(shoken: str, mondai: list[dict], model: str = '') -> tuple[str, dict]:
    kiroku = {'書き直しを依頼した文': len(mondai), '書き直せた文': 0, '削除した文': 0}
    if not mondai:
        return shoken, kiroku
    try:
        from agent_sdk import llm_call
    except ImportError:
        return shoken, kiroku

    bunmen = '\n'.join(f'{i}. {x["文"].strip()}\n   問題: {x["理由"]}'
                       for i, x in enumerate(mondai, 1))
    prompt = (
        '債務者区分の業況分析所見の一部に不備があります。**渡した文だけを直してください。**\n'
        '他の文は渡していません。文の意味を変えず、書かれていない事実を足さないこと。\n\n'
        '【タグの決まり】文末にタグを1つだけ付ける。文中や文頭に付けてはならない。\n'
        '  【事実】＝カルテに記載がある内容・カルテの数値をそのまま述べる文。評価の語を含めない\n'
        '  【推測】＝カルテの記載から推定できるが確認が取れていない内容\n'
        '  【一般論】＝外部資料に基づく業種一般の傾向。★外部資料由来は必ずこれ\n'
        '  【意見】＝評価・判断・推奨を含む文（「留意が必要」「懸念される」等）\n\n'
        '【資料の説明にとどまる文の扱い】\n'
        '  「本資料は〜の調査である」「調査事項には〜が含まれる」のように、'
        'その資料が何を調べているかを述べただけの文は所見に載せない。\n'
        '  業界がどうなっているか（数値・水準・方向）を述べる文に書き直す。\n'
        '  ★渡された文の中にその材料が無いなら、書き足さずに「削除」の1語だけを返す。\n\n'
        f'【直す文】\n{bunmen}\n\n'
        '同じ順番・同じ件数で返してください。'
    )
    try:
        r = await llm_call(prompt=prompt, schema=_SHUSEI_SCHEMA, model=model or LLM_MODEL)
    except Exception as e:
        _glog(f'所見の部分修正に失敗しました（元の文を残します）: {type(e).__name__}: {e}')
        return shoken, kiroku

    naoshi = [str(x.get('修正後の文') or '').strip()
              for x in ((r.get('data') or {}).get('修正') or [])]
    if len(naoshi) != len(mondai):
        _glog(f'所見の部分修正の件数が合いません（依頼{len(mondai)}／返り{len(naoshi)}）。元の文を残します')
        return shoken, kiroku

    for x, atarashii in zip(mondai, naoshi):
        if not atarashii:
            continue
        if atarashii == '削除':
            shoken = shoken.replace(x['文'], '')
            kiroku['削除した文'] += 1
        else:
            shoken = shoken.replace(x['文'], atarashii if atarashii.endswith('。')
                                    else atarashii + '。')
            kiroku['書き直せた文'] += 1
    return re.sub(r'\n{3,}', '\n\n', shoken).strip(), kiroku

def _nen_mikeisai_chushaku(shoken: str, zenbun: list, gaibu_no: set) -> tuple[str, dict]:
    nen_nashi = [y for y in (zenbun or []) if not y.get('nen')]
    inyou = [y for y in nen_nashi if y['no'] in gaibu_no]
    meiki = NEN_MEIKI_GO in (shoken or '')
    kensa = {'年未掲載の採用': len(nen_nashi), 'うち引用された': len(inyou),
             '明記あり': meiki, 'コードで追記': False}
    if not inyou or meiki:
        return '', kensa
    kensa['コードで追記'] = True
    namae = '／'.join(f'「{str(y.get("title"))[:44]}」' for y in inyou)
    return f'※{namae}は年次が資料に記載されていません。鮮度は確認できていません。', kensa


def _tsukeru_kotowari(shoken: str, gaibu_ari: bool, chushaku: str = '') -> str:
    if not gaibu_ari:
        return shoken
    if GYOKAI_KOTOWARI in shoken:
        return shoken
    atama = GYOKAI_KOTOWARI
    if chushaku:
        atama += '\n' + chushaku
    return f'{atama}\n{shoken}'


async def write_output_excel(
    template_path: str,
    output_path: str,
    extracted: dict,
    llm_tokurei_hanteibun: str,
    llm_saishu_kubun: str,
    llm_hantei_konkyo: str,
    llm_kaisha_gaiyo_yoyaku: str,
    gyokyo_bunseki: dict[str, dict],
    llm_gyokai_shoken: str,
    gyokai_no_inyou: set,
    used_model: str = '(モデル名未取得)',
    gyokai_joho: dict | None = None,
) -> dict:
    if extracted.get('status') != 'ok':
        raise Exception('extracted のstatusが"ok"ではありません。不足情報がある状態で出力しようとしています。')

    template_path = _resolve_template_path(template_path)
    shutil.copy(template_path, output_path)
    wb = openpyxl.load_workbook(output_path)
    ws = wb['分析シート']

    gaiyo = extracted['gaiyo']
    meta = extracted['meta']
    hantei = extracted['hantei']

    ws[OUTPUT_CELLS['対象先名']] = gaiyo.get('商号')
    ws[OUTPUT_CELLS['顧客番号']] = gaiyo.get('顧客番号')
    ws[OUTPUT_CELLS['判定基準日']] = datetime.now(JST).strftime('%Y/%m/%d')
    ws[OUTPUT_CELLS['直近決算期']] = meta.get('直近決算期')
    ws[OUTPUT_CELLS['部店名']] = meta.get('部店')
    ws[OUTPUT_CELLS['担当者']] = meta.get('担当者')
    ws[OUTPUT_CELLS['会社概要要約']] = llm_kaisha_gaiyo_yoyaku

    ws[OUTPUT_CELLS['縦軸']] = f"{hantei['採用行']}　{hantei['採用行名']}"
    ws[OUTPUT_CELLS['横軸']] = f"{hantei['横軸記号']}　{hantei['横軸名']}"
    ws[OUTPUT_CELLS['交点']] = hantei['交点区分']
    ws[OUTPUT_CELLS['特例']] = llm_tokurei_hanteibun
    ws[OUTPUT_CELLS['最終区分']] = llm_saishu_kubun
    ws[OUTPUT_CELLS['判定根拠']] = llm_hantei_konkyo
    _z = (gyokai_joho or {}).get('zenbun')
    chushaku, nen_kensa = _nen_mikeisai_chushaku(llm_gyokai_shoken, _z, gyokai_no_inyou)
    if chushaku:
        _glog(f'年次未掲載の明記漏れを検知したため注記を足しました: {chushaku}')
    if _GAIBU_TRACE:
        _GAIBU_TRACE.setdefault('所見の検査', {}).update(nen_kensa)
    llm_gyokai_shoken = _tsukeru_kotowari(llm_gyokai_shoken, bool(gyokai_no_inyou), chushaku)
    ws[OUTPUT_CELLS['業界特異性所見']] = llm_gyokai_shoken

    from openpyxl.styles import Alignment

    for shihyo_mei, row in GYOKYO_BUNSEKI_ROWS.items():
        vals = gyokyo_bunseki.get(shihyo_mei)
        if not vals:
            logger.warning(f'業況分析の指標「{shihyo_mei}」の値が渡されていません')
            continue
        ws.cell(row, 3, vals.get('前々期'))
        ws.cell(row, 4, vals.get('前期'))
        ws.cell(row, 5, vals.get('当期'))
        ws.cell(row, 6, vals.get('前期比'))
        ws.cell(row, 7, vals.get('フラグ'))
        ws.cell(row, 8, vals.get('注記'))
        ws.cell(row, 9, vals.get('確認事項'))

        numfmt = NUMFMT_SHOSU if shihyo_mei in SHOSU_ITEMS else NUMFMT_KINGAKU
        for col in (3, 4, 5):
            cell = ws.cell(row, col)
            if isinstance(cell.value, (int, float)):
                cell.number_format = numfmt
            cell.alignment = Alignment(horizontal='right', vertical='center')
        ws.cell(row, 6).alignment = Alignment(horizontal='right', vertical='center')

        ws.row_dimensions[row].height = None

    orphan = {name: v['確認事項'] for name, v in gyokyo_bunseki.items()
              if v.get('確認事項') and name not in GYOKYO_BUNSEKI_ROWS}
    if orphan:
        logger.info(f'様式に出力行がないため確認事項をExcelに書けなかった指標: {orphan}')

    chishiki = ['債務者区分判定実施基準 別表1（形式区分マトリックス）／注①〜⑫（2026.4改定・原本PDF確認済み）']
    if gyokai_joho and gyokai_joho.get('zenbun') is not None:
        z = gyokai_joho['zenbun']
        kensu = len(gyokai_joho.get('queries') or [gyokai_joho.get('query')])
        chishiki.append(f"【外部リサーチ】検索{kensu}本"
                        f"／異なるURL{len(gyokai_joho.get('results') or [])}件・採用{len(z)}件")
        chishiki.append('  検索クエリ: '
                        + '／'.join(gyokai_joho.get('queries')
                                    or [str(gyokai_joho.get('query'))]))
        inyou = [y for y in z if y['no'] in gyokai_no_inyou]
        if inyou:
            chishiki.append(f'■ 全文を確認して採用し、所見で引用した資料（{len(inyou)}件）')
            for y in inyou:
                chishiki.append(f"  {y['no']}. {y['title']}")
                chishiki.append(
                    f"     対象業種={y['対象業種']}／対象地域={y['対象地域']}"
                    f"／対象企業規模={y['対象企業規模']}")
                chishiki.append(
                    '     ' + (f"データの年次={y['データの年次']}" if y.get('nen')
                              else NEN_MIKEISAI)
                    + '／' + ('固定で毎回取得' if y.get('固定') else
                             f"{kensu}本の検索のうち{y.get('出現本数', 1)}本で出現")
                    + ('／一覧から深掘り' if y.get('深掘り') else ''))
                kijutsu = str(y.get('当社業種の記述') or '').strip()
                if kijutsu:
                    chishiki.append(f"     引用できた記述: {kijutsu[:90]}")
                chishiki.append(f"     {y['url']}")
        elif z:
            chishiki.append(f'  {len(z)}件を採用しましたが、所見では引用されていません')
        else:
            chishiki.append('  採用できる資料がありませんでした')
        mi = [y for y in z if y not in inyou]
        if mi:
            logger.info('採用したが所見で引用されなかった資料: '
                        + '／'.join(str(y.get('title'))[:40] for y in mi))
    elif gyokai_joho and gyokai_joho.get('results'):
        chishiki.append(f"【外部リサーチ・業界動向】検索クエリ「{gyokai_joho.get('query')}」")
        for i, x in enumerate(gyokai_joho['results'], 1):
            chishiki.append(f"  {i}. {x['title']} {x['url']}")
    ws[OUTPUT_CELLS['参照ナレッジ']] = '\n'.join(chishiki)
    ws[OUTPUT_CELLS['参照カルテ']] = f"{gaiyo.get('商号')}様カルテ"
    ws[OUTPUT_CELLS['使用モデル生成日時']] = f"{used_model}／{datetime.now(JST).isoformat()}"

    wb.save(output_path)
    logger.info(f'分析シートを出力しました: {output_path}')
    return {'output_path': output_path, '最終区分': llm_saishu_kubun}


_ZOKA_GA_AKKA_ITEMS = {
    '流動負債', '実質債務', '有利子負債合計', '売上債権回転期間', '買入債務回転期間',
    '実質債務償還年数', '金融費用率',
    '短期借入金', '長期借入金', 'その他有利子負債（リース割賦他）', '支払利息',
}
_GENSHO_GA_AKKA_ITEMS = {
    '売上高', '営業利益', '経常利益', '税引後利益', '総資産', '自己資本', '流動資産',
    '自己資本比率', '流動比率', '正常運転資金', '繰越利益剰余金',
}

KAKUNIN_AKKA_RITSU = 0.05
KAKUNIN_MAX_KENSU = 5

_KAKUNIN_HOKO = {
    '買入債務回転期間': 'up', '売上債権回転期間': 'up', '棚卸資産': 'up',
    '実質債務償還年数': 'up', '金融費用率': 'up',
    '自己資本比率': 'down', '流動比率': 'down', '自己資本': 'down',
    '資本金': 'down', '営業利益': 'down',
}

_KAKUNIN_YUSEN = {
    '実質債務償還年数': 1, '自己資本': 1, '営業利益': 1,
    '買入債務回転期間': 2, '資本金': 2,
    '流動比率': 3, '自己資本比率': 3, '金融費用率': 3,
    '売上債権回転期間': 4, '棚卸資産': 4,
}

_KAKUNIN_KAITEN_KIKAN_KA = {'買入債務回転期間': 0.5, '売上債権回転期間': 0.5}


_KAKUNIN_SUIJUN = {
    '自己資本': lambda cur: cur < 0,
    '実質債務償還年数': lambda cur: cur > 15,
}


def _kin(v) -> str:
    return '不明' if v is None else f'{v:,.0f}'


def _sho(v) -> str:
    return '不明' if v is None else f'{v:,.2f}'.rstrip('0').rstrip('.')


def _kakunin_bun_shokan(p1: float, cur: float) -> str:
    if cur > 15:
        kyori = '15年を超えており縦軸6に該当する'
    elif cur > 10:
        kyori = f'15年（縦軸6）まで残り{15 - cur:.1f}年'
    else:
        kyori = f'10年（縦軸5）まで残り{10 - cur:.1f}年'
    if p1 is None or cur > p1:
        return (f'返済能力の低下（{_sho(p1)}年→{_sho(cur)}年）。{kyori}。'
                f'借入増と収益減のどちらが要因か、改善計画の有無を確認')
    return (f'前期からは改善（{_sho(p1)}年→{_sho(cur)}年）。ただし{kyori}。'
            f'改善が一過性か（資産売却・特別利益等）を確かめ、償還年数の見通しを確認')


def _kakunin_bun_jikoshihon(p1: float | None, cur: float) -> str:
    if cur >= 0:
        return (f'自己資本の減少（{_kin(p1)}→{_kin(cur)}百万円）。損失による毀損の進行度と、'
                f'債務超過までの余裕、増資・役員借入の予定を確認')
    if p1 is not None and p1 >= 0:
        return (f'債務超過に転落（{_kin(p1)}→{_kin(cur)}百万円）。解消の見通し（増資・利益計画）と、'
                f'代表者等からの借入金を自己資本にみなせるか（注③）を確認')
    return (f'債務超過が継続（{_kin(p1)}→{_kin(cur)}百万円）。解消までの年数と改善計画、'
            f'代表者等からの借入金を自己資本にみなせるか（注③）を確認')


def _kakunin_bun_ryudo(p1: float, cur: float) -> str:
    warigome = '。100%を割っており短期の支払いに手元資金の取り崩しが必要な状態' if cur < 100 else ''
    return (f'短期の支払能力の低下（{_sho(p1)}%→{_sho(cur)}%）{warigome}。'
            f'手元資金の月商倍率と当座の資金繰り予定を確認')


def _kakunin_bun_saiken(p1: float, cur: float, uriage_zoka: bool) -> str:
    if uriage_zoka:
        return (f'回収期間の長期化（{_sho(p1)}月→{_sho(cur)}月）。売上は増加しているため、'
                f'期末に売上が集中していないか、計上時期と入金予定を確認')
    return (f'回収期間の長期化（{_sho(p1)}月→{_sho(cur)}月）。売上は減少しており回収遅延の疑いがある。'
            f'滞留している売掛金・受取手形の有無と、大口先の支払条件変更を確認')


def _kakunin_bun_tanaoroshi(p1: float, cur: float, uriage_zoka: bool) -> str:
    if uriage_zoka:
        return (f'在庫の増加（{_kin(p1)}→{_kin(cur)}百万円）。売上も増加しているため繁忙期に向けた'
                f'仕込みの可能性がある。在庫の内訳と回転状況を確認')
    return (f'在庫の増加（{_kin(p1)}→{_kin(cur)}百万円）。売上が減る中での増加であり滞留の疑いがある。'
            f'不良在庫の有無と評価方法、実地棚卸の状況を確認')


_KAKUNIN_BUN = {
    '買入債務回転期間': lambda p1, cur, _: (
        f'支払サイトの伸び（{_sho(p1)}月→{_sho(cur)}月）。仕入先への支払いが遅れていないか、'
        f'支払条件の変更を申し入れていないかを確認'),
    '実質債務償還年数': lambda p1, cur, _: _kakunin_bun_shokan(p1, cur),
    '流動比率': lambda p1, cur, _: _kakunin_bun_ryudo(p1, cur),
    '売上債権回転期間': _kakunin_bun_saiken,
    '棚卸資産': _kakunin_bun_tanaoroshi,
    '自己資本比率': lambda p1, cur, _: (
        f'財務基盤の毀損（{_sho(p1)}%→{_sho(cur)}%）。損失によるものか資産の増加によるものかを分け、'
        f'増資・役員借入による補強の予定を確認'),
    '金融費用率': lambda p1, cur, _: (
        f'借入負担の増加（{_sho(p1)}%→{_sho(cur)}%）。借入増と金利上昇のどちらが要因か、'
        f'新規借入の使途と返済計画を確認'),
    '自己資本': lambda p1, cur, _: _kakunin_bun_jikoshihon(p1, cur),
    '資本金': lambda p1, cur, _: (
        f'減資（{_kin(p1)}→{_kin(cur)}百万円）。欠損填補・税負担軽減等の目的と、'
        f'登記の有無・稟議の経緯を確認'),
    '営業利益': lambda p1, cur, _: (
        f'本業の採算悪化（{_kin(p1)}→{_kin(cur)}百万円）。不採算案件の有無、価格転嫁の状況、'
        f'売上の増減との関係を確認'),
}


def _kakunin_akka_ritsu(indicator: str, p1: float | None, cur: float | None) -> float | None:
    hoko = _KAKUNIN_HOKO.get(indicator)
    if hoko is None or p1 is None or cur is None:
        return None
    akka = cur > p1 if hoko == 'up' else cur < p1
    if not akka:
        return None
    ka = _KAKUNIN_KAITEN_KIKAN_KA.get(indicator)
    if ka is not None and max(abs(p1), abs(cur)) < ka:
        return None
    if p1 == 0:
        return float('inf')
    return abs((cur - p1) / p1)


def _apply_kakunin_jiko(result: dict[str, dict]) -> None:
    uriage = result.get('売上高') or {}
    uriage_zoka = (uriage.get('前期') is not None and uriage.get('当期') is not None
                   and uriage['当期'] > uriage['前期'])

    kesson = [name for name, v in result.items() if v.get('確認事項')]

    kouho = []
    for indicator, v in result.items():
        if v.get('確認事項'):
            continue
        cur = v.get('当期')

        suijun = _KAKUNIN_SUIJUN.get(indicator)
        if suijun is not None and cur is not None and suijun(cur):
            bun = _KAKUNIN_BUN[indicator](v['前期'], cur, uriage_zoka)
            kouho.append((_KAKUNIN_YUSEN[indicator], float('-inf'), indicator, bun))
            continue

        ritsu = _kakunin_akka_ritsu(indicator, v.get('前期'), cur)
        if ritsu is None or ritsu < KAKUNIN_AKKA_RITSU:
            continue
        bun = _KAKUNIN_BUN[indicator](v['前期'], cur, uriage_zoka)
        kouho.append((_KAKUNIN_YUSEN[indicator], -ritsu, indicator, bun))

    kouho.sort(key=lambda x: (x[0], x[1]))
    saiyo = kouho[:KAKUNIN_MAX_KENSU]
    hoshutsu = kouho[KAKUNIN_MAX_KENSU:]

    for _, order, indicator, bun in saiyo:
        v = result[indicator]
        v['確認事項'] = bun
        if not v['フラグ']:
            v['フラグ'] = FLAG_YOKAKUNIN_KAKUNIN
        elif v['フラグ'] == FLAG_KAIZEN:
            v['フラグ'] = FLAG_YOKAKUNIN_KAKUNIN
        wa_suijun = order == float('-inf')
        if not wa_suijun and not any(k in (v['注記'] or '') for k in ('悪化', '改善')):
            v['注記'] = ((v['注記'] + '／') if v['注記'] else '') + (
                f"悪化方向（{_fmt_hyoji(indicator, v['前期'])}→"
                f"{_fmt_hyoji(indicator, v['当期'])}）")
        if indicator not in GYOKYO_BUNSEKI_ROWS:
            logger.info(f'確認事項を作ったが様式に出力行がない指標: {indicator}')

    if hoshutsu:
        logger.info(
            f'確認事項は優先順位上位{KAKUNIN_MAX_KENSU}件に絞った。'
            f'枠外（フラグは残る）: ' + '、'.join(f'{n}(優先{y})' for y, _, n, _ in hoshutsu)
        )
    logger.info(
        '確認事項の出力（悪化・水準）: ' + ('、'.join(n for _, _, n, _ in saiyo) or 'なし')
        + '／（算出不能・データ欠損）: ' + ('、'.join(kesson) or 'なし')
    )


FLAG_KAIZEN = '【改善】良い方向に大きく変化'
FLAG_YOKAKUNIN_10 = '【要確認】±10%以上の変化'
FLAG_YOKAKUNIN_KAKUNIN = '【要確認】確認事項あり'
FLAG_YOKAKUNIN_ZERO = '【要確認】前期が0で変化率を出せない'
FLAG_DAIHUKU = '【大幅増減】±20%以上の変化'
FLAG_TREND = '【トレンド悪化】3期連続で悪化方向'
FLAG_SANTEI_FUNO = '【算出不能】数値を出せない'
FLAG_KESSON = '【データ欠損】数値が無い'
FLAGS_HENKA = (FLAG_YOKAKUNIN_10, FLAG_DAIHUKU)


def _fmt_hyoji(indicator: str, v: float | None) -> str:
    if v is None:
        return '算出不能'
    if indicator in SHOSU_ITEMS:
        return f'{v:,.2f}'.rstrip('0').rstrip('.')
    return f'{v:,.0f}'


def _flag_for_period(prev: float | None, curr: float | None) -> tuple[str, str]:
    if prev is None or curr is None:
        return '', FLAG_KESSON
    if prev == 0:
        if curr == 0:
            return '±0%', ''
        return '算出不可（前期0）', FLAG_YOKAKUNIN_ZERO
    ch = (curr - prev) / abs(prev)
    if abs(ch) >= 0.20:
        return f'{ch:+.1%}', FLAG_DAIHUKU
    if abs(ch) >= 0.10:
        return f'{ch:+.1%}', FLAG_YOKAKUNIN_10
    return f'{ch:+.1%}', ''


def _change_direction(indicator: str, p1: float | None, cur: float | None) -> str:
    if p1 is None or cur is None or p1 == cur:
        return ''
    increased = cur > p1
    if indicator in _ZOKA_GA_AKKA_ITEMS:
        return '悪化方向' if increased else '改善方向'
    if indicator in _GENSHO_GA_AKKA_ITEMS:
        return '改善方向' if increased else '悪化方向'
    hoko = _KAKUNIN_HOKO.get(indicator)
    if hoko == 'up':
        return '悪化方向' if increased else '改善方向'
    if hoko == 'down':
        return '改善方向' if increased else '悪化方向'
    return ''


def _is_trend_worsening(indicator: str, p2: float | None, p1: float | None, cur: float | None) -> bool:
    if None in (p2, p1, cur):
        return False
    if indicator in _ZOKA_GA_AKKA_ITEMS:
        return p2 < p1 < cur
    if indicator in _GENSHO_GA_AKKA_ITEMS:
        return p2 > p1 > cur
    return False


def _safe_div(numerator: float | None, denominator: float | None,
              multiplier: float = 1.0) -> float | None:
    if numerator is None or denominator is None or denominator == 0:
        return None
    return round(numerator / denominator * multiplier, 2)


def _derive_indicators(zaimu: dict[str, dict[str, float | None]]) -> dict[str, dict[str, float | None]]:
    periods = ('前々期', '前期', '当期')
    names = ['実質債務', '自己資本比率', '流動比率', '売上債権回転期間',
             '買入債務回転期間', '正常運転資金', '実質債務償還年数', '金融費用率']
    derived: dict[str, dict[str, float | None]] = {n: {} for n in names}

    for p in periods:
        def g(item: str) -> float | None:
            return zaimu.get(item, {}).get(p)

        uketori, urikake, tanaoroshi = g('受取手形'), g('売掛金'), g('棚卸資産')
        ryudo_shiharai, kaikake = g('流動性支払手形'), g('買掛金')
        uriage = g('売上高')

        unten_parts = (uketori, urikake, tanaoroshi, ryudo_shiharai, kaikake)
        unten = (uketori + urikake + tanaoroshi - ryudo_shiharai - kaikake
                 if all(v is not None for v in unten_parts) else None)
        derived['正常運転資金'][p] = unten

        yurishi = g('有利子負債合計')
        if yurishi is not None and unten is not None:
            jisshitsu_saimu = yurishi - unten
        else:
            jisshitsu_saimu = None
        derived['実質債務'][p] = jisshitsu_saimu

        zeibiki, genka = g('税引後利益'), g('減価償却費')
        bunbo = zeibiki + genka if None not in (zeibiki, genka) else None
        derived['実質債務償還年数'][p] = (
            _safe_div(jisshitsu_saimu, bunbo) if bunbo is not None and bunbo > 0 else None
        )

        derived['自己資本比率'][p] = _safe_div(g('自己資本'), g('総資産'), 100)
        derived['流動比率'][p] = _safe_div(g('流動資産'), g('流動負債'), 100)

        gessho = g('月商')
        if gessho is None and uriage is not None:
            gessho = uriage / 12
        waribiki = g('割引・譲渡手形') or 0
        uriage_saiken = (uketori + urikake + waribiki
                         if None not in (uketori, urikake) else None)
        kainyu_saimu = (ryudo_shiharai + kaikake if None not in (ryudo_shiharai, kaikake) else None)
        derived['売上債権回転期間'][p] = _safe_div(uriage_saiken, gessho)
        derived['買入債務回転期間'][p] = _safe_div(kainyu_saimu, gessho)

        derived['金融費用率'][p] = _safe_div(g('支払利息'), g('売上総利益'), 100)

    return derived


SEARCH_INCLUDE_DOMAINS = [
    'go.jp',
    'or.jp',
    'tdb.co.jp',
    'tsr-net.co.jp',
    'meti.go.jp',
    'mlit.go.jp',
    'boj.or.jp',
    'jfc.go.jp',
    'chusho.meti.go.jp',
]
SEARCH_MAX_RESULTS = 8
SEARCH_PER_DOMAIN_MAX = 3
SEARCH_CONTENT_CHARS = 1500
SEARCH_MIN_CONTENT = 200
SEARCH_JUTSUGO_LIST = ('統計', '調査', '景気動向', '')
SEARCH_JUTSUGO = SEARCH_JUTSUGO_LIST[0]
SEARCH_SHUGO_MAX = 2
SEARCH_QUERY_MAX = 8
SEARCH_POOL_MAX = 30
FUKABORI_SHIKII = 1
FUKABORI_MAP_DEPTH = 1
FUKABORI_MAP_LIMIT = 50
FUKABORI_PDF_MAX = 5

ZENBUN_HANTEI = True
JITSU_TEXT_MIN = 500
ZENBUN_AI_MAX = 10
ZENBUN_PROMPT_CHARS = 15000
NEN_MIKEISAI = '年が未掲載のため鮮度は不明。参考として掲載'
NEN_MEIKI_GO = '年次未掲載'
NEN_KAGEN = 3
SAIYO_MAX = 3
GYOSHU_ITCHI_NG = ('不一致',)
GYOSHU_ITCHI_YUSEN = ('一致', '類似', '判定不能')
GYOSHU_KIJUTSU_NASHI = ('なし', '不明', '')


def _url_seiki(url: str) -> str:
    u = re.sub(r'^https?://', '', str(url or ''))
    u = u.split('#')[0].split('?')[0]
    u = re.sub(r'/{2,}', '/', u)
    return u.rstrip('/').lower()


GAIBU_LOG_PRINT = True
GAIBU_JSON_PRINT = False


def _glog(msg: str) -> None:
    logger.info(msg)
    if GAIBU_LOG_PRINT:
        print(msg)


def _build_search_shugo(gyoshu: str | None, toriatsukai: list | None) -> list[str]:
    if not gyoshu:
        return []
    gyoshu = str(gyoshu).strip()
    hinmoku = [str(x).strip() for x in (toriatsukai or []) if str(x).strip()]

    kouho = []
    if gyoshu.startswith('その他') or len(gyoshu) <= 3:
        if hinmoku:
            kouho.append(hinmoku[0])
        nokori = gyoshu.replace('その他', '', 1).strip()
        if nokori:
            kouho.append(nokori)
    else:
        kouho.append(gyoshu)
        if hinmoku and hinmoku[0] != gyoshu:
            kouho.append(hinmoku[0])

    out = []
    for k in kouho:
        if k and k not in out:
            out.append(k)
    return out[:SEARCH_SHUGO_MAX]


def _build_search_queries(gyoshu: str | None, toriatsukai: list | None) -> list[str]:
    shugo = _build_search_shugo(gyoshu, toriatsukai)
    if not shugo:
        return []
    out = []
    for s in shugo:
        for j in SEARCH_JUTSUGO_LIST:
            q = f'{s} {j}'.strip()
            if q not in out:
                out.append(q)
    return out[:SEARCH_QUERY_MAX]


def _build_search_query(gyoshu: str | None, toriatsukai: list | None) -> str | None:
    q = _build_search_queries(gyoshu, toriatsukai)
    return q[0] if q else None


_GAIBU_TRACE: dict = {}


def _trace_reset() -> None:
    _GAIBU_TRACE.clear()
    _GAIBU_TRACE.update({'許可ドメイン': list(SEARCH_INCLUDE_DOMAINS),
                         '検索': [], 'プール': [], '全文判定': {}, '深掘り': {}})


def _gaibu_json(gyokai_joho: dict | None, shogo: str = '',
                output_dir: str = '') -> str | None:
    if not _GAIBU_TRACE:
        return None
    g = gyokai_joho or {}
    _GAIBU_TRACE['採用'] = g.get('zenbun') or []
    _GAIBU_TRACE['コスト'] = {
        'tavilyクレジット': g.get('credits', 0),
        '抽出モデル': ZENBUN_MODEL,
        **(g.get('zenbun_kosuto') or {}),
    }
    body = json.dumps(_GAIBU_TRACE, ensure_ascii=False, indent=2)
    if GAIBU_JSON_PRINT:
        print('━━ 外部リサーチログ（JSON）ここから')
        print(body)
        print('━━ 外部リサーチログ（JSON）ここまで')
    else:
        z = _GAIBU_TRACE.get('全文判定') or {}
        k = _GAIBU_TRACE.get('コスト') or {}
        print(f'━━ 外部リサーチ　検索{len(_GAIBU_TRACE.get("検索") or [])}本'
              f'／取得{z.get("取得", 0)}件 → 採用{z.get("採用", 0)}件'
              f'／{k.get("tavilyクレジット", 0)}クレジット'
              f'／llm_call {k.get("llm_call", 0)}回')
    try:
        name = _FILENAME_NG.sub('_', str(shogo or '判定').strip()) or '判定'
        path = os.path.join(output_dir or OUTPUT_DIR, f'外部リサーチログ_{name}.json')
        if os.path.dirname(path):
            os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, 'w', encoding='utf-8') as f:
            f.write(body)
        _glog(f'外部リサーチログを書き出しました: {path}')
        return path
    except Exception as e:
        _glog(f'外部リサーチログの書き出しに失敗しました（処理は続行）: {type(e).__name__}: {e}')
        return None


KOTEI_SHIRYO_ON = True
KOTEI_KOHO_TSUKI = 4
KOTEI_SHIRYO = (
    {'name': '日本銀行 企業向けサービス価格指数',
     'url': 'https://www.boj.or.jp/statistics/pi/cspi_release/sppi{yy}{mm}.pdf',
     'note': '月次。大類別に リース・レンタル／不動産／運輸・郵便／情報通信／'
             '諸サービス（労働者派遣・土木建築・建物）／広告／金融・保険 を持つ'},
)


def _kotei_urls(kijun: datetime | None = None) -> list[tuple[str, str]]:
    kijun = kijun or datetime.now(JST)
    out = []
    for s in KOTEI_SHIRYO:
        y, mo = kijun.year, kijun.month
        for _ in range(KOTEI_KOHO_TSUKI):
            mo -= 1
            if mo == 0:
                y, mo = y - 1, 12
            out.append((s['url'].format(yy=f'{y % 100:02d}', mm=f'{mo:02d}'), s['name']))
    return out


async def _kotei_shutoku(kijun: datetime | None = None) -> list[dict]:
    if not KOTEI_SHIRYO_ON:
        return []
    try:
        from agent_sdk import tavily_extract
    except ImportError:
        _glog('tavily_extractが無い環境のため固定資料の取得を省略します')
        return []

    kouho = _kotei_urls(kijun)
    try:
        r = await tavily_extract(urls=[u for u, _ in kouho], extract_depth='advanced')
    except Exception as e:
        _glog(f'固定資料の取得に失敗しました（判定は続行します）: {type(e).__name__}: {e}')
        return []

    torete = {str(x.get('url') or ''): x for x in (r.get('results') or [])}
    out = []
    for u, name in kouho:
        x = torete.get(u)
        if x is None:
            continue
        kouzou = await _pdf_kouzou_chushutsu(u)
        if kouzou:
            honbun, jp = kouzou
        else:
            honbun, jp = _jitsu_text(x.get('raw_content'))
        _glog(f'  固定資料を取得: {name}（日本語{jp}字／表構造={bool(kouzou)}）{u}')
        if jp < JITSU_TEXT_MIN:
            _glog(f'    実テキスト{jp}字で使えないため見送ります')
            continue
        out.append({
            'title': name,
            'url': u,
            'content': honbun[:SEARCH_CONTENT_CHARS],
            'raw_content': x.get('raw_content'),
            '出現本数': 0,
            'クエリ': [],
            '固定': True,
        })
        break
    if not out:
        _glog('  固定資料は取得できませんでした')
    _GAIBU_TRACE['固定資料'] = [
        {'url': u, '資料名': name, '取得': u in torete} for u, name in kouho]
    return out

async def search_gyokai_joho(gyoshu: str | None, toriatsukai: list | None = None) -> dict:
    queries = _build_search_queries(gyoshu, toriatsukai)
    if not queries:
        _glog('業種が取得できていないため業界情報の検索を行いません')
        return {'query': None, 'queries': [], 'results': [], 'credits': 0}

    _trace_reset()
    rs = await asyncio.gather(*[_search_one(q) for q in queries])
    for q, r in zip(queries, rs):
        _GAIBU_TRACE['検索'].append({
            'クエリ': q, '取得件数': len(r['results']),
            'クレジット': r['credits'],
            'urls': [{'url': x['url'], 'title': x['title'],
                      '抜粋字数': len(str(x.get('content') or '')),
                      '全文字数': len(str(x.get('raw_content') or ''))}
                     for x in r['results']],
        })

    pool: dict[str, dict] = {}
    for q, r in zip(queries, rs):
        for x in r['results']:
            k = _url_seiki(x['url'])
            e = pool.get(k)
            if e is None:
                e = {**x, '出現本数': 0, 'クエリ': []}
                pool[k] = e
            if q not in e['クエリ']:
                e['出現本数'] += 1
                e['クエリ'].append(q)
            if len(str(x.get('content') or '')) > len(str(e.get('content') or '')):
                e['content'] = x['content']
            if len(str(x.get('raw_content') or '')) > len(str(e.get('raw_content') or '')):
                e['raw_content'] = x['raw_content']

    results = sorted(pool.values(),
                     key=lambda z: (-z['出現本数'], -len(str(z.get('content') or ''))))
    results = results[:SEARCH_POOL_MAX]
    kotei = await _kotei_shutoku()
    kotei = [k for k in kotei if _url_seiki(k['url']) not in pool]
    results = kotei + results
    credits = sum(r['credits'] for r in rs)

    _glog(f'━━ 外部リサーチ　クエリ{len(queries)}本／{credits}クレジット'
          f'／検索で異なるURL {len(pool)}件'
          + (f'／固定資料 {len(kotei)}件' if kotei else '／固定資料なし'))
    _glog(f'   許可ドメイン: {"／".join(SEARCH_INCLUDE_DOMAINS)}')
    for q, r in zip(queries, rs):
        _glog(f'  ◆「{q}」 取得{len(r["results"])}件')
        for x in r['results']:
            _glog(f'      {x["url"]}')
        if not r['results']:
            _glog('      （0件）')
    fukusu = [z for z in results if z['出現本数'] >= 2]
    _glog(f'  ── 束ねた結果　2本以上に出てきたURL {len(fukusu)}件')
    for z in results[:10]:
        _glog(f'    {"固定" if z.get("固定") else str(z["出現本数"]) + "本"}'
              f' {str(z["title"])[:44]}')
        _glog(f'         {z["url"]}')
    if not results:
        _glog('  採用できる業界情報がありませんでした。この情報なしで続行します')
    _GAIBU_TRACE['プール'] = [
        {'no': n, 'url': z['url'], 'title': z['title'],
         '出現本数': z['出現本数'], '出現クエリ': z['クエリ']}
        for n, z in enumerate(results, 1)]
    return {'query': queries[0], 'queries': queries,
            'results': results, 'credits': credits}


async def _search_one(query: str) -> dict:
    try:
        from agent_sdk import tavily_search
    except ImportError:
        _glog('tavily_searchが利用できない環境のため業界情報の検索を省略します')
        return {'query': query, 'results': [], 'credits': 0}

    try:
        r = await tavily_search(query=query, max_results=SEARCH_MAX_RESULTS,
                                include_domains=SEARCH_INCLUDE_DOMAINS,
                                include_raw_content=True)
    except Exception as e:
        _glog(f'検索に失敗しました（判定は続行します）「{query}」: {type(e).__name__}: {e}')
        return {'query': query, 'results': [], 'credits': 0}

    results, seen = [], set()
    for x in (r.get('results') or []):
        url = str(x.get('url') or '')
        body = str(x.get('content') or '')
        if len(body) < SEARCH_MIN_CONTENT:
            continue
        seiki = _url_seiki(url)
        if seiki in seen:
            continue
        seen.add(seiki)
        results.append({
            'title': str(x.get('title') or '')[:120],
            'url': url,
            'content': body[:SEARCH_CONTENT_CHARS],
            'raw_content': x.get('raw_content'),
        })
    return {'query': query, 'results': results,
            'credits': (r.get('usage') or {}).get('credits', 0)}


_B64 = re.compile(r'data:[a-zA-Z/+.-]+;base64,[A-Za-z0-9+/=\s]{200,}|[A-Za-z0-9+/=]{300,}')
_JP = re.compile(r'[ぁ-んァ-ヶ一-龥、。「」（）％]')
_N4 = re.compile(r'(?<![0-9])(19[89][0-9]|20[0-9]{2})(?![0-9])')
_YMD = re.compile(r'(?<![0-9])(20[0-9]{2})(?:0[1-9]|1[0-2])(?:[0-2][0-9]|3[01])(?![0-9])')
_YM = re.compile(r'(?<![0-9])(20[0-9]{2})(?:0[1-9]|1[0-2])(?![0-9])')
_WAREKI = re.compile(r'(令和|平成)\s*([0-9]{1,2}|元)')
_HD = re.compile(r'/h([0-9]{1,2})/')

_ZENBUN_SCHEMA_KEYS = ('発行年', 'データの年次', '年の根拠', '対象業種',
                       '対象地域', '対象企業規模', '業種の一致', '当社業種の記述', '要点')


def _jitsu_text(raw: str) -> tuple[str, int]:
    h = _B64.sub(' ', str(raw or ''))
    return h, len(_JP.findall(h))


def _nen_kouho(s: str) -> list[int]:
    s = s or ''
    o = [int(x) for x in _N4.findall(s)]
    o += [int(x) for x in _YMD.findall(s)] + [int(x) for x in _YM.findall(s)]
    for gengo, n in _WAREKI.findall(s):
        v = 1 if n == '元' else int(n)
        o.append((2018 + v) if gengo == '令和' else (1988 + v))
    o += [1988 + int(m.group(1)) for m in _HD.finditer(s)]
    return sorted({y for y in o if 1985 <= y <= 2100})


def _to_nen(v) -> int | None:
    m = _N4.search(str(v or ''))
    return int(m.group()) if m else None


ZENBUN_ATAMA_CHARS = 1500
KIRIDASHI_MAE = 800
KIRIDASHI_ATO = 3000
KIRIDASHI_MAX = 3


def _kiridashi_go(gyoshu: str, gyokai_joho: dict | None) -> list[str]:
    go, mita = [], set()
    for q in ((gyokai_joho or {}).get('queries') or []):
        s = q
        for j in SEARCH_JUTSUGO_LIST:
            if j and s.endswith(' ' + j):
                s = s[:-(len(j) + 1)]
        s = s.strip()
        if s and s not in mita:
            mita.add(s)
            go.append(s)
    if gyoshu and gyoshu not in mita:
        go.append(gyoshu)
    return sorted(go, key=len, reverse=True)


def _honbun_kiridashi(honbun: str, go: list[str]) -> tuple[str, list[int]]:
    if len(honbun) <= ZENBUN_PROMPT_CHARS:
        return honbun, []

    atama = honbun[:ZENBUN_ATAMA_CHARS]
    mado, atta = [], []
    for g in go:
        for m in re.finditer(re.escape(g), honbun):
            i = m.start()
            if i < ZENBUN_ATAMA_CHARS:
                continue
            if any(s <= i <= e for s, e in mado):
                continue
            mado.append((max(0, i - KIRIDASHI_MAE), min(len(honbun), i + KIRIDASHI_ATO)))
            atta.append(i)
            if len(mado) >= KIRIDASHI_MAX:
                break
        if len(mado) >= KIRIDASHI_MAX:
            break

    if not mado:
        return honbun[:ZENBUN_PROMPT_CHARS], []

    mado.sort()
    bubun = [atama]
    for s, e in mado:
        bubun.append(f'\n……（{s}字目から）……\n' + honbun[s:e])
    out = ''.join(bubun)
    return out[:ZENBUN_PROMPT_CHARS], atta

def _zenbun_schema(gyoshu: str) -> dict:
    return {
        'type': 'object',
        'properties': {
            '発行年': {'type': 'string', 'description':
                     '資料が公表・更新された年（西暦4桁）。和暦は西暦に直す。'
                     '★本文に根拠がなければ推測せず「不明」'},
            'データの年次': {'type': 'string', 'description':
                        '★中身のデータが何年のものか（西暦4桁）。発行年とは別。'
                        '例「平成11年度調査（平成10年4月期〜平成11年3月期決算）」なら1998。'
                        '読み取れなければ「不明」'},
            '年の根拠': {'type': 'string', 'description':
                     '発行年とデータの年次をどこから読み取ったか本文を引用。'
                     '引用できないなら「引用できない」'},
            '対象業種': {'type': 'string', 'description': '資料が対象としている業種。なければ「記載なし」'},
            '対象地域': {'type': 'string', 'description': '全国／四国地区など。なければ「記載なし」'},
            '対象企業規模': {'type': 'string', 'description':
                       '大手53社／従業者50人未満など。なければ「記載なし」'},
            '業種の一致': {'type': 'string', 'enum': ['一致', '類似', '不一致', '判定不能'],
                      'description': f'当社の業種「{gyoshu}」と資料の対象業種の関係。'
                                     f'日本標準産業分類で分類が異なるなら「不一致」'},
            '当社業種の記述': {'type': 'string', 'description':
                        f'★当社の業種について述べている箇所を本文からそのまま1つ引用する。'
                        f'資料が対象業種を宣言していなくても、当社の業種に触れていれば引用する。'
                        f'本文に【構造化した表】の節がある場合は、その中の当社業種に該当する'
                        f'1行をそのまま引用してよい（ラベルと数値の対応が保たれている）。'
                        f'★【構造化した表】の節が無い本文では、表の行や数字の羅列を引用しない'
                        f'（改行が失われ、どの数字がどの項目のものか判別できないことがある）。'
                        f'引用できる箇所が無ければ「なし」。★探して無ければ作らないこと'},
            '要点': {'type': 'string', 'description': '当社の業況分析に使える内容を3行以内。なければ「なし」'},
        },
        'required': list(_ZENBUN_SCHEMA_KEYS),
    }


_ZENBUN_KIHON = {
    '発行年': '不明', 'データの年次': '不明', '年の根拠': '引用できない',
    '対象業種': '記載なし', '対象地域': '記載なし', '対象企業規模': '記載なし',
    '業種の一致': '判定不能', '当社業種の記述': 'なし', '要点': 'なし',
}


def _zenbun_umeru(no: int, data: dict | None) -> dict:
    d = dict(data or {})
    kake = [k for k in _ZENBUN_SCHEMA_KEYS if not str(d.get(k) or '').strip()]
    for k in kake:
        d[k] = _ZENBUN_KIHON[k]
    if kake:
        _glog(f'  {no}. 抽出結果に欠けている項目がありました（既定値で埋めます）: {kake}')
    return {k: d[k] for k in _ZENBUN_SCHEMA_KEYS}


async def _zenbun_chushutsu(no: int, x: dict, honbun: str,
                            gyoshu: str, chiiki: str, schema: dict,
                            kiridashi_go: list[str] | None = None) -> dict:
    try:
        from agent_sdk import llm_call
    except ImportError:
        return {'no': no, 'err': 'llm_callが利用できない環境'}
    watasu_honbun, atta = _honbun_kiridashi(honbun, kiridashi_go or [])
    if atta:
        _glog(f'  {no}. 全文{len(honbun):,}字のうち、業種語の周辺を抜き出して渡します'
              f'（{atta}字目付近）')
    prompt = (
        f'以下はWeb上の資料の全文です。当社（業種「{gyoshu}」・所在地「{chiiki}」）の'
        f'業況分析に使えるかを人が判断するため、**事実だけを抜き出してください**。'
        f'評価や意見は書かないこと。\n'
        f'★「発行年」と「データの年次」を必ず分けてください。資料が2019年に公表されていても、'
        f'中身が1998年度のデータなら データの年次は1998 です。'
        f'与信判断ではデータの年次が重要です。\n'
        f'★本文に根拠がないものは「不明」「記載なし」と書いてください。**推測で埋めないこと。**\n'
        f'　年は本文のどこから読み取ったかを「年の根拠」に引用してください。\n'
        f'★「当社業種の記述」には、当社の業種について述べている箇所を引用してください。\n'
        f'　本文に【構造化した表】の節がある場合は、その中の該当行をそのまま引用してよい\n'
        f'　（ラベルと数値の対応が保たれています）。\n'
        f'　★【構造化した表】の節が無い本文では、表の行や数字の羅列を引用しないでください。\n'
        f'　改行が失われており、どの数字がどの項目のものか判別できず、誤った値を引くことに\n'
        f'　なります。文になっている箇所が無ければ「なし」と書いてください。\n'
        f'　「要点」も同様に、【構造化した表】に基づく場合だけ数値を書いてよい。\n'
        f'【タイトル】{x.get("title")}\n【URL】{x.get("url")}\n'
        f'【全文】\n{watasu_honbun}'
    )
    try:
        r = await llm_call(prompt=prompt, schema=schema, model=ZENBUN_MODEL)
    except Exception as e:
        return {'no': no, 'err': f'{type(e).__name__}: {e}'}
    usage = r.get('usage') or {}
    return {'no': no, 'data': _zenbun_umeru(no, r.get('data')),
            'in': usage.get('input_tokens') or 0, 'out': usage.get('output_tokens') or 0}


def _otosu(otoshita: list, no, x, riyu: str) -> None:
    x = x or {}
    otoshita.append((no, x.get('title'), riyu, str(x.get('url') or '')))


FUKABORI_JOGAI_DOMAIN = ('tdb.co.jp', 'tsr-net.co.jp')


def _fukabori_taisho(url: str) -> bool:
    dom = _url_seiki(url).split('/')[0]
    if url.lower().endswith('.pdf'):
        return True
    if any(dom.endswith(d) for d in FUKABORI_JOGAI_DOMAIN):
        return False
    return dom.endswith('go.jp') or dom.endswith('or.jp')


async def _fukabori(yomeru: list, saiyo_urls: set, gyoshu: str, chiiki: str,
                    schema: dict, kiri_go: list[str] | None = None
                    ) -> tuple[list, list, dict]:
    try:
        from agent_sdk import tavily_map, tavily_extract
    except ImportError:
        _glog('tavily_map／tavily_extractが無い環境のため深掘りを省略します')
        return [], [], {'llm_call': 0, 'in': 0, 'out': 0}

    nokori = [(i, x, h) for i, x, h in yomeru if str(x.get('url')) not in saiyo_urls]
    if not nokori:
        return [], [], {'llm_call': 0, 'in': 0, 'out': 0}
    nokori.sort(key=lambda z: -z[1].get('出現本数', 1))
    oya_no, oya, _ = nokori[0]
    _glog(f'深掘り：採用が{FUKABORI_SHIKII}件以下のため'
                f'{oya.get("出現本数", 1)}本に出現したページを1回だけ掘ります'
                f'｜{str(oya.get("title"))[:44]}')

    try:
        m = await tavily_map(url=oya['url'], max_depth=FUKABORI_MAP_DEPTH,
                             limit=FUKABORI_MAP_LIMIT)
    except TypeError:
        try:
            m = await tavily_map(url=oya['url'])
        except Exception as e:
            _glog(f'深掘りのtavily_mapに失敗しました: {type(e).__name__}: {e}')
            return [], [], {'llm_call': 0, 'in': 0, 'out': 0}
    except Exception as e:
        _glog(f'深掘りのtavily_mapに失敗しました: {type(e).__name__}: {e}')
        return [], [], {'llm_call': 0, 'in': 0, 'out': 0}

    seen, pdf = set(), []
    for u in (m.get('results') or []):
        u = str(u)
        if not _fukabori_taisho(u):
            continue
        k = _url_seiki(u)
        if k in seen:
            continue
        seen.add(k)
        pdf.append(u)
    pdf = pdf[:FUKABORI_PDF_MAX]
    _glog(f'  子URL {len(m.get("results") or [])}件 → 対象 {len(pdf)}件')
    if not pdf:
        _glog('  読める見込みの子が無いため深掘りを打ち切ります')
        fuka_oto = []
        _otosu(fuka_oto, oya_no, oya, '深掘りしたがPDFの子が無い')
        return [], fuka_oto, {'llm_call': 0, 'in': 0, 'out': 0}

    try:
        ex = await tavily_extract(urls=pdf, extract_depth='advanced')
    except Exception as e:
        _glog(f'深掘りのtavily_extractに失敗しました: {type(e).__name__}: {e}')
        return [], [], {'llm_call': 0, 'in': 0, 'out': 0}

    watasu, otoshita = [], []
    for n, r in enumerate(ex.get('results') or [], 1):
        kouzou = await _pdf_kouzou_chushutsu(str(r.get('url') or ''))
        if kouzou:
            honbun, jp = kouzou
        else:
            honbun, jp = _jitsu_text(r.get('raw_content'))
        title = str(r.get('url') or '').rsplit('/', 1)[-1]
        if jp < JITSU_TEXT_MIN:
            _otosu(otoshita, f'深{n}', {'title': title, 'url': r.get('url')},
                   f'実テキスト{jp}字で判読不能')
            continue
        watasu.append((f'深{n}', {'title': title, 'url': str(r.get('url') or ''),
                                  '出現本数': oya.get('出現本数', 1),
                                  'クエリ': oya.get('クエリ') or []}, honbun))
    for fr in (ex.get('failed_results') or []):
        _otosu(otoshita, '深', {'title': str(fr.get('url') or '').rsplit('/', 1)[-1],
                                'url': fr.get('url')}, '取得失敗')
    if not watasu:
        return [], otoshita, {'llm_call': 0, 'in': 0, 'out': 0}

    js = await asyncio.gather(*[
        _zenbun_chushutsu(i, x, h, gyoshu, chiiki, schema, kiri_go) for i, x, h in watasu])
    saiyo = []
    for j, (i, x, _) in zip(js, watasu):
        if 'err' in j:
            _otosu(otoshita, i, x, f'抽出失敗（{j["err"]}）')
            continue
        d = j['data']
        nen = _to_nen(d['データの年次']) or _to_nen(d['発行年'])
        if nen and KONNEN - nen > NEN_KAGEN:
            _otosu(otoshita, i, x, f'{nen}年＝{KONNEN - nen}年前')
            continue
        if d['業種の一致'] in GYOSHU_ITCHI_NG:
            _otosu(otoshita, i, x, f'業種が不一致（{d["対象業種"]}）')
            continue
        if str(d['当社業種の記述']).strip() in GYOSHU_KIJUTSU_NASHI:
            _otosu(otoshita, i, x, '当社の業種の記述が引用できない')
            continue
        if str(d['要点']).strip() in ('なし', ''):
            _otosu(otoshita, i, x, '使える要点がない')
            continue
        saiyo.append({'no': i, 'nen': nen, 'title': x['title'], 'url': x['url'],
                      '出現本数': x['出現本数'], '出現クエリ': x['クエリ'],
                      '深掘り': True, **{k: d[k] for k in _ZENBUN_SCHEMA_KEYS}})
    kosuto = {'llm_call': len(watasu),
              'in': sum(j.get('in', 0) for j in js),
              'out': sum(j.get('out', 0) for j in js)}
    _glog(f'  深掘りの結果：採用{len(saiyo)}件／除外{len(otoshita)}件'
                f'／llm_call {kosuto["llm_call"]}回')
    return saiyo, otoshita, kosuto


def _pdf_gyo_seiri(rows: list) -> list[str]:
    out = []
    for row in rows:
        cells = [str(c).replace('\n', ' ').strip() for c in row if c and str(c).strip()]
        if len(cells) >= 2:
            out.append('  '.join(cells))
    return out


async def _pdf_kouzou_chushutsu(url: str) -> tuple[str, int] | None:
    if not url.lower().endswith('.pdf'):
        return None
    try:
        import pdfplumber
        import urllib.request
        import io
    except ImportError:
        return None
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        data = urllib.request.urlopen(req, timeout=20).read()
    except Exception as e:
        _glog(f'  PDFの直接取得に失敗: {type(e).__name__}: {e}')
        return None
    # 表の検出方法は2種類あり、どちらが当たるかはページごとに違う。
    # 境界線のある表は線ベースでしか拾えない（テキスト位置ベースだと崩れて0件になった）。
    # 境界線の無い表（位置揃えだけの表）は線ベースでは拾えない（実測：セルに複数行が
    # 押し込まれて崩れた）。件数や行数で優劣を決めても質は測れないため、両方使う。
    try:
        gyo, mongon = [], []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                page_gyo = []
                for senryaku in (None, {'vertical_strategy': 'text', 'horizontal_strategy': 'text'}):
                    for t in (page.extract_tables(table_settings=senryaku) if senryaku
                             else page.extract_tables()) or []:
                        page_gyo += _pdf_gyo_seiri(t)
                txt = page.extract_text() or ''
                if page_gyo:
                    # 表が何の業種・項目の表かは、表のセルではなく同じページの見出し文に
                    # 書かれていることがある（実測：「建設機械器具賃貸業」は本文側にあり、
                    # 表の行には出てこなかった）。見出しを表の直前に添えて対応関係を保つ。
                    midashi = re.sub(r'\s+', ' ', txt[:200]).strip()
                    if midashi:
                        gyo.append(f'［{midashi}］')
                    gyo += page_gyo
                if txt:
                    mongon.append(txt)
    except Exception as e:
        _glog(f'  PDFの表構造抽出に失敗: {type(e).__name__}: {e}')
        return None
    if not gyo:
        return None
    seen, gyo_yuni = set(), []
    for g in gyo:
        if g not in seen:
            seen.add(g)
            gyo_yuni.append(g)
    honbun = ('【本文（先頭ページ）】\n' + (mongon[0][:2000] if mongon else '')
              + '\n\n【構造化した表（1行が1レコード。ラベルと数値の対応関係を保っている。'
                '地の文中の数字より、この行の数字を優先すること）】\n'
              + '\n'.join(gyo_yuni))
    _, jp = _jitsu_text(honbun)
    return honbun, jp


async def zenbun_hantei(gyokai_joho: dict, gyoshu: str, chiiki: str) -> dict:
    results = (gyokai_joho or {}).get('results') or []
    if not results:
        return {**(gyokai_joho or {}), 'zenbun': [], 'otoshita': [],
                'zenbun_kosuto': {'llm_call': 0, 'in': 0, 'out': 0}}

    yomeru, otoshita, meisai = [], [], {}
    for i, x in enumerate(results, 1):
        kouzou = await _pdf_kouzou_chushutsu(str(x.get('url') or ''))
        if kouzou:
            honbun, jp = kouzou
            x = {**x, '表構造': True}
            _glog(f'  {i}. PDFを直接取得して表構造を保った本文に差し替えました（日本語{jp}字）')
        else:
            honbun, jp = _jitsu_text(x.get('raw_content'))
        meisai[i] = {'no': i, 'url': x['url'], 'title': x['title'],
                     '出現本数': x.get('出現本数', 1), '日本語字数': jp,
                     '結果': '', '理由': '', 'AI抽出': None}
        if jp < JITSU_TEXT_MIN:
            _otosu(otoshita, i, x, f'実テキスト{jp}字で判読不能')
            continue
        yomeru.append((i, x, honbun))

    watasu, dom_count = [], {}
    for i, x, honbun in yomeru:
        if x.get('固定'):
            watasu.append((i, x, honbun))
    for i, x, honbun in yomeru:
        if x.get('固定'):
            continue
        dom = _url_seiki(x['url']).split('/')[0]
        if dom_count.get(dom, 0) >= SEARCH_PER_DOMAIN_MAX:
            _otosu(otoshita, i, x, f'同一ドメイン{SEARCH_PER_DOMAIN_MAX}件超（{dom}）')
            continue
        if len(watasu) >= ZENBUN_AI_MAX:
            _otosu(otoshita, i, x, f'AIに渡す上限{ZENBUN_AI_MAX}件超')
            continue
        dom_count[dom] = dom_count.get(dom, 0) + 1
        watasu.append((i, x, honbun))

    _glog(f'全文判定：取得{len(results)}件 → 読める{len(yomeru)}件'
                f' → AIに渡す{len(watasu)}件')
    for i, x, _ in watasu:
        moto = '固定資料' if x.get('固定') else f'{x.get("出現本数", 1)}本に出現'
        _glog(f'  AIに渡す {moto}｜{str(x.get("title"))[:44]}')
    if not watasu:
        _glog('全文を読める資料がありませんでした。業界情報なしで続行します')
        return {**gyokai_joho, 'zenbun': [], 'otoshita': otoshita,
                'zenbun_kosuto': {'llm_call': 0, 'in': 0, 'out': 0}}

    schema = _zenbun_schema(gyoshu)
    kiri_go = _kiridashi_go(gyoshu, gyokai_joho)
    js = await asyncio.gather(*[
        _zenbun_chushutsu(i, x, h, gyoshu, chiiki, schema, kiri_go) for i, x, h in watasu])

    saiyo = []
    for j, (i, x, _) in zip(js, watasu):
        if 'err' in j:
            _otosu(otoshita, i, x, f'抽出失敗（{j["err"]}）')
            continue
        d = j['data']
        meisai[i]['AI抽出'] = dict(d)
        meisai[i]['トークン'] = {'in': j.get('in', 0), 'out': j.get('out', 0)}
        nen = _to_nen(d['データの年次']) or _to_nen(d['発行年'])
        kouho = _nen_kouho(f'{x.get("url")} {x.get("title")}')
        if nen and kouho and nen not in kouho:
            _glog(f'  {i}. 年の不一致 AI={nen}／URL・タイトル={kouho}'
                        f'（AIの申告を採用）｜{str(x.get("title"))[:40]}')
        if nen and KONNEN - nen > NEN_KAGEN:
            _otosu(otoshita, i, x, f'{nen}年＝{KONNEN - nen}年前')
            continue
        if d['業種の一致'] in GYOSHU_ITCHI_NG:
            _otosu(otoshita, i, x, f'業種が不一致（{d["対象業種"]}）')
            continue
        if str(d['当社業種の記述']).strip() in GYOSHU_KIJUTSU_NASHI:
            _otosu(otoshita, i, x, '当社の業種の記述が引用できない')
            continue
        if str(d['要点']).strip() in ('なし', ''):
            _otosu(otoshita, i, x, '使える要点がない')
            continue
        saiyo.append({'no': i, 'nen': nen,
                      'title': str(x.get('title') or ''), 'url': str(x.get('url') or ''),
                      '出現本数': x.get('出現本数', 1),
                      '出現クエリ': x.get('クエリ') or [],
                      '固定': bool(x.get('固定')),
                      **{k: d[k] for k in _ZENBUN_SCHEMA_KEYS}})

    fuka_kosuto = {'llm_call': 0, 'in': 0, 'out': 0}
    if len(saiyo) <= FUKABORI_SHIKII:
        tsuika, fuka_otoshita, fuka_kosuto = await _fukabori(
            yomeru, {z['url'] for z in saiyo}, gyoshu, chiiki, schema, kiri_go)
        saiyo += tsuika
        otoshita += fuka_otoshita
        _GAIBU_TRACE['深掘り'] = {
            '走った': True, '採用': len(tsuika), 'コスト': fuka_kosuto,
            '採用したもの': [{'url': z['url'], 'title': z['title'],
                        'データの年次': z['データの年次']} for z in tsuika],
            '落としたもの': [{'no': str(n), 'title': str(t), '理由': r, 'url': u}
                       for n, t, r, u in fuka_otoshita],
        }
    else:
        _GAIBU_TRACE['深掘り'] = {'走った': False,
                              '理由': f'採用が{len(saiyo)}件あり閾値{FUKABORI_SHIKII}件を超えている'}

    for no, t, riyu, _u in otoshita:
        if isinstance(no, int) and no in meisai:
            meisai[no]['結果'], meisai[no]['理由'] = '除外', riyu
    for z in saiyo:
        if z['no'] in meisai:
            meisai[z['no']]['結果'] = '採用'
    _GAIBU_TRACE['全文判定'] = {
        '取得': len(results), '読める': len(yomeru), 'AIに渡した': len(watasu),
        '採用': len(saiyo),
        '条件': {'実テキスト下限': JITSU_TEXT_MIN, '何年前まで': NEN_KAGEN,
               '業種で落とす': list(GYOSHU_ITCHI_NG),
               '当社業種の記述が必須': True,
               'AIに渡す上限': ZENBUN_AI_MAX, '同一ドメイン上限': SEARCH_PER_DOMAIN_MAX,
               '所見に渡す上限': SAIYO_MAX},
        '明細': [meisai[k] for k in sorted(meisai)],
    }

    saiyo.sort(key=lambda z: (
        GYOSHU_ITCHI_YUSEN.index(z['業種の一致']) if z['業種の一致'] in GYOSHU_ITCHI_YUSEN
        else len(GYOSHU_ITCHI_YUSEN),
        0 if z.get('固定') else 1,
        0 if z['nen'] else 1,
        -z['出現本数'], -(z['nen'] or 0)))
    wakugai = saiyo[SAIYO_MAX:]
    saiyo = saiyo[:SAIYO_MAX]
    kosuto = {'llm_call': len(watasu) + fuka_kosuto['llm_call'],
              'in': sum(j.get('in', 0) for j in js) + fuka_kosuto['in'],
              'out': sum(j.get('out', 0) for j in js) + fuka_kosuto['out']}

    _glog(f'全文判定：採用{len(saiyo)}件／除外{len(otoshita)}件'
                f'／llm_call {kosuto["llm_call"]}回'
                f'（in {kosuto["in"]:,} out {kosuto["out"]:,}）')
    for z in saiyo:
        _glog(f'  採用 {z["no"]}. {str(z["nen"]) + "年" if z["nen"] else "年未掲載"}'
              f' [{z["業種の一致"]}] {z["title"][:44]}')
    for i, t, riyu, u in otoshita:
        _glog(f'  除外 {i}. {riyu}｜{str(t)[:44]}｜{u[:70]}')
    if wakugai:
        _glog(f'  枠外（採用上限{SAIYO_MAX}件超）{len(wakugai)}件: '
              + '／'.join(f'{z["nen"] or "年未掲載"} {z["title"][:30]}' for z in wakugai))
    if not saiyo:
        _glog('採用できる資料がありませんでした。業界情報なしで続行します')
    return {**gyokai_joho, 'zenbun': saiyo, 'otoshita': otoshita,
            'zenbun_kosuto': kosuto}


def _calculate_gyokyo_bunseki(zaimu: dict[str, dict[str, float | None]],
                              karte_shihyo: dict[str, dict[str, float | None]] | None = None
                              ) -> dict[str, dict]:
    derived = _derive_indicators(zaimu)
    if karte_shihyo:
        for name, periods in karte_shihyo.items():
            base = dict(derived.get(name, {}))
            for period, value in periods.items():
                if value is not None:
                    base[period] = value
            derived[name] = base

    PERIODS = ('前々期', '前期', '当期')
    santei_funo: dict[str, dict[str, str]] = {}
    for period in PERIODS:
        zeibiki = zaimu.get('税引後利益', {}).get(period)
        genka = zaimu.get('減価償却費', {}).get(period)
        yurishi = zaimu.get('有利子負債合計', {}).get(period)
        unten = derived.get('正常運転資金', {}).get(period)
        bunbo = None if None in (zeibiki, genka) else zeibiki + genka
        jisshitsu = None if None in (yurishi, unten) else yurishi - unten

        if jisshitsu is not None and jisshitsu <= 0:
            santei_funo[period] = {
                '注記': (f'正常運転資金{unten:,.0f}が有利子負債{yurishi:,.0f}を上回り、'
                       f'実質債務は{jisshitsu:,.0f}（実質無借金）のため算出対象外'),
                '確認事項': '',
            }
        elif bunbo is not None and bunbo <= 0:
            santei_funo[period] = {
                '注記': (f'税引後利益{zeibiki:,.0f}＋減価償却費{genka:,.0f}＝{bunbo:,.0f}で'
                       f'分母が0以下のため算出不能'),
                '確認事項': '赤字により返済能力が測れない。改善計画と資金繰りの見通しを確認',
            }

    if '実質債務償還年数' in derived:
        y = dict(derived['実質債務償還年数'])
        for period in PERIODS:
            if y.get(period) == 0:
                y[period] = None
                if period not in santei_funo:
                    santei_funo[period] = {
                        '注記': 'カルテに0と記載されており算出できなかったものとして扱う'
                                '（有利子負債がある先で償還年数0年はあり得ないため）',
                        '確認事項': 'カルテ上の償還年数が0となっている理由を確認',
                    }
        derived['実質債務償還年数'] = y

    merged: dict[str, dict[str, float | None]] = {**zaimu, **derived}

    result: dict[str, dict] = {}
    for indicator, periods in merged.items():
        p2, p1, cur = periods.get('前々期'), periods.get('前期'), periods.get('当期')
        change_str, flag = _flag_for_period(p1, cur)
        note = ''
        kakunin = ''

        if flag in FLAGS_HENKA and p1 is not None and cur is not None:
            direction = _change_direction(indicator, p1, cur)
            if direction == '改善方向':
                flag = FLAG_KAIZEN
                note = ((note + '／') if note else '') + \
                    f'改善方向（{_fmt_hyoji(indicator, p1)}→{_fmt_hyoji(indicator, cur)}）'
            elif direction:
                note = ((note + '／') if note else '') + \
                    f'{direction}（{_fmt_hyoji(indicator, p1)}→{_fmt_hyoji(indicator, cur)}）'
            else:
                note = ((note + '／') if note else '') + \
                    '増減の良し悪しを一律に判断できない指標のため方向は付けない'

        if _is_trend_worsening(indicator, p2, p1, cur):
            flag = FLAG_TREND if not flag else flag
            note = ((note + '／') if note else '') + (
                f'3期連続で悪化方向（{_fmt_hyoji(indicator, p2)}→'
                f'{_fmt_hyoji(indicator, p1)}→{_fmt_hyoji(indicator, cur)}）')

        if indicator == '実質債務' and cur is not None and cur < 0:
            flag = ''
            note = f'正常運転資金が有利子負債を上回っており実質無借金の状態（実質債務{cur:,.0f}）'
            kakunin = ''

        if indicator == '実質債務償還年数' and cur is not None and cur < 0:
            jisshitsu_cur = merged.get('実質債務', {}).get('当期')
            if jisshitsu_cur is not None and jisshitsu_cur < 0:
                flag = ''
                note = ('正常運転資金が有利子負債を上回っており実質無借金の状態'
                        '（マイナス値のため償還年数の評価対象外）')
                kakunin = ''
            else:
                flag = FLAG_SANTEI_FUNO
                note = santei_funo.get('当期', {}).get('注記') or (
                    f'償還年数が{cur}とマイナスだが実質債務は'
                    f'{"不明" if jisshitsu_cur is None else format(jisshitsu_cur, ",.0f")}で'
                    f'マイナスではない。分母（税引後利益＋減価償却費）がマイナスのため算出不能'
                )
                kakunin = santei_funo.get('当期', {}).get('確認事項') or (
                    '赤字により返済能力が測れない。改善計画と資金繰りの見通しを確認'
                )
                logger.info(
                    f'実質債務償還年数がマイナス（{cur}）だが実質債務は{jisshitsu_cur}。'
                    f'実質無借金ではなく分母マイナスによる算出不能として扱う'
                )

        if indicator == '実質債務償還年数' and cur is None and '当期' in santei_funo:
            flag = FLAG_SANTEI_FUNO
            note = santei_funo['当期']['注記']
            kakunin = santei_funo['当期']['確認事項']
        elif indicator == '実質債務' and cur == 0 and '当期' in santei_funo:
            note = (note + '／' if note else '') + santei_funo['当期']['注記']
        elif flag == FLAG_KESSON:
            kakunin = kakunin or '当期の数値が未取得。カルテの該当欄を確認'

        result[indicator] = {
            '前々期': p2, '前期': p1, '当期': cur,
            '前期比': change_str, 'フラグ': flag, '注記': note, '確認事項': kakunin,
        }

    _apply_kakunin_jiko(result)
    return result


_LLM_RESPONSE_SCHEMA = {
    'type': 'object',
    'properties': {
        '特例適用判断': {
            'type': 'string',
            'description': (
                '特例（注②③⑥⑦⑧⑩⑪⑫）の適用可否とその理由。'
                '適用条件を抽出データから確認できない場合は「適用の余地あり」に留め、断定しない。'
                '★注①（低位区分優先の判定ルール）・注④（償還年数の算式）・'
                '注⑤（個人事業者の行選択）・注⑨（横軸Bの定義）は特例ではないため、'
                'この欄に「注④の特例を適用し…」のように書いてはならない。'
                '★交点に括弧書きの代替区分がない場合は「特例適用なし」と書くこと。'
            ),
        },
        '最終形式区分': {'type': 'string', 'enum': ['正常先', '正常先2', '要注意先', '要管理先', '破綻懸念先', '実質破綻先', '破綻先']},
        '判定根拠': {
            'type': 'string',
            'description': (
                '「・」で始まる3〜5点の箇条書き。各点に必ず具体的な数値根拠を含める。'
                '★箇条書きの各行にタグ（【事実】【意見】等）を付けてはならない'
                '（タグを付けるのは業界特異性所見だけ）。'
                '★信頼度表示は文章全体のいちばん最後に1回だけ付すこと。'
                '各行の末尾に繰り返し付けてはならない。'
                '信頼度表示は次のいずれか1つを"そのままの表記で"使う：'
                '特例適用なし かつ カルテに欠損なしの場合は「【社内基準に基づく形式判定】」、'
                '特例を適用した または カルテに欠損がある場合は「【推測を含む形式判定】」。'
                'これ以外の信頼度表記（「信頼度:高」等）を独自に作ってはならない。'
                '★複数の縦軸に該当して低位区分優先（注①）が適用された場合は、'
                'どの縦軸に該当しなぜその行を採用したかを必ず1点として含めること。'
            ),
        },
        '会社概要要約': {
            'type': 'string',
            'description': (
                '150字±10文字（140〜160字）。次の5要素をすべて必ず含める：'
                '①業種・取扱品目 ②設立年・沿革の要点 ③上場区分・資本金・資本構成 '
                '④事業所展開・代表者 ⑤直近売上規模。'
                '特に④の事業所展開（拠点・従業員数など）の記載漏れが起きやすいので必ず入れる。'
                '「業界大手」「安定成長」等の未検証の形容は使わない。数値は原文表記（百万円）を維持する。'
            ),
        },
        '業界特異性所見': {
            'type': 'array',
            'description': (
                'カルテ由来の着眼点。配列の1要素が1文にあたる。2〜4個。'
                '★タグ（【事実】等）は文中に書かない。文の本文だけを「文」に入れ、'
                'タグは「種類」から機械的に付けるため、AIが書く必要はない。'
                '★1要素に事実と評価を混ぜてはならない。数値や事象を述べる要素と、'
                'それに対する評価の要素を必ず分け、事実の要素を先に置く。'
                '★「事実」の種類を1個以上必ず含めること。'
                'カルテの数値に触れているのに「事実」が1つもないのは不可。'
            ),
            'items': {
                'type': 'object',
                'properties': {
                    '種類': {'type': 'string', 'enum': ['事実', '推測', '意見'],
                            'description':
                                '事実＝カルテの数値をそのまま述べる。評価の語を含めない。'
                                '推測＝カルテの記載から推定できるが確認が取れていない内容。'
                                '意見＝評価・判断・推奨を含む文'
                                '（「留意が必要」「懸念される」等の語を含む文は意見）。'},
                    '文': {'type': 'string',
                          'description': 'タグを付けない、文の本文だけ。「。」で終える。'},
                },
                'required': ['種類', '文'],
            },
        },
    },
    'required': ['特例適用判断', '最終形式区分', '判定根拠', '会社概要要約', '業界特異性所見'],
}


SHINRAIDO_KEISHIKI = '【社内基準に基づく形式判定】'
SHINRAIDO_SUISOKU = '【推測を含む形式判定】'

GAIYO_MOJISU = 150
GAIYO_MOJISU_KYOYO = 10

KONKYO_MAX_KENSU = 5


def _expected_shinraido(kouten: str, kesson: bool) -> str:
    return SHINRAIDO_SUISOKU if ('（' in kouten or kesson) else SHINRAIDO_KEISHIKI


def _normalize_hantei_konkyo(text: str, kouten: str, kesson: bool) -> str:
    if not text:
        return text
    before = text

    for tag in (SHINRAIDO_KEISHIKI, SHINRAIDO_SUISOKU):
        text = text.replace(tag, '')
    text = text.rstrip('　 \n')

    text = re.sub(r'(?<!\A)(?<!\n)・', '\n・', text).rstrip('　 \n')

    lines = text.split('\n')
    kept, dropped, count = [], [], 0
    for line in lines:
        if line.lstrip().startswith('・'):
            count += 1
            if count > KONKYO_MAX_KENSU:
                dropped.append(line)
                continue
        elif dropped:
            dropped.append(line)
            continue
        kept.append(line)
    if dropped:
        logger.info(
            f'判定根拠が{count}点あったため上限{KONKYO_MAX_KENSU}点に切った。'
            f'枠外: ' + '／'.join(x.strip()[:40] for x in dropped if x.strip())
        )
    text = '\n'.join(kept).rstrip('　 \n')

    expected = _expected_shinraido(kouten, kesson)
    text = f'{text}{expected}'
    if text != before:
        logger.info('判定根拠の体裁を整えた（改行・信頼度表示）')
    return text


def _check_gaiyo_mojisu(gaiyo_yoyaku: str) -> str:
    n = len(gaiyo_yoyaku or '')
    lo, hi = GAIYO_MOJISU - GAIYO_MOJISU_KYOYO, GAIYO_MOJISU + GAIYO_MOJISU_KYOYO
    if lo <= n <= hi:
        return ''
    msg = f'会社概要要約が{n}字（規定{GAIYO_MOJISU}字±{GAIYO_MOJISU_KYOYO}字の範囲外）'
    logger.info(msg)
    return msg


def _build_llm_prompt(extracted: dict, gyokyo_bunseki: dict[str, dict]) -> str:
    hantei = extracted['hantei']
    gaiyo = extracted['gaiyo']
    keisan = extracted['keisan']
    zaimu_cur = {k: v.get('当期') for k, v in extracted['zaimu'].items()}

    kouten_str = hantei['交点区分']

    uchiwake = hantei.get('縦軸判定内訳') or {}
    jujiku_uchiwake = '／'.join(
        f'{name}={"該当" if uchiwake.get(name) else "非該当"}' for name in ('債務超過', '赤字', '繰損')
    ) if uchiwake else '（内訳なし）'
    if uchiwake.get('繰損') and not uchiwake.get('赤字'):
        jujiku_uchiwake += (
            '　★当期は黒字だが繰損があるため縦軸1（黒字かつ繰損なし）には該当しない。'
            'この点を判定根拠に明記すること'
        )

    daitai_note = ''
    if '（' in kouten_str:
        daitai_note = (
            f'交点は「{kouten_str}」で、括弧内は代替区分の候補です。'
            f'特例（注{TOKUREI_NO_BANGO}）のうち該当しうる条項の条件を、'
            f'抽出データから確認できる範囲で判断してください。'
            f'条件を満たすと確認できない場合は「適用の余地あり」に留め、断定しないでください。'
            f'★この場合「特例適用なし」と書き出してはいけません。'
            f'「なし」と「余地あり」を同じ文に並べると、読み手はどちらなのか判断できません。'
            f'「注○の適用余地あり：（理由）」または「注○の適用条件を満たさない：（理由）」の'
            f'いずれかの形で書いてください。'
        )
    else:
        daitai_note = (
            f'交点「{kouten_str}」には括弧書きの代替区分がありません。'
            f'この場合「特例適用なし。交点「{kouten_str}」には括弧書きの代替区分がないため、'
            f'適用しうる特例がない」のように、理由まで書くこと。'
            f'「特例適用なし」の4文字だけで終えてはならない。'
            f'存在しない特例を適用したかのように書かないこと。'
        )

    teii_note = ''
    if hantei['低位区分優先を適用したか']:
        kouho = hantei['該当行候補']
        meisai = '、'.join(
            f'縦軸{n}（{ROW_NAMES[n]}）→ {MATRIX_TABLE[n][hantei["横軸記号"]]}' for n in kouho
        )
        teii_note = (
            f'\n【★判定根拠に必ず含めること：低位区分優先（注①）の適用】\n'
            f'この先は複数の縦軸に該当しています：{meisai}。\n'
            f'注①により低位（悪い方）の縦軸{hantei["採用行"]}を採用しました。\n'
            f'判定根拠には「どの縦軸に該当したか」「なぜその行を採用したか（注①）」を'
            f'必ず1点として明記してください。これを省略すると、審査担当者は'
            f'なぜこの区分になったのか判断できません。'
        )

    def _fmt(v: object) -> str:
        return '算出不能' if v is None else str(v)

    def _num(v: object) -> str:
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return '（記載なし）' if v is None else str(v)

    flagged_lines = []
    for name, v in gyokyo_bunseki.items():
        if not v['フラグ']:
            continue
        line = (f"- {name}: {_fmt_hyoji(name, v['前々期'])} → {_fmt_hyoji(name, v['前期'])}"
                f" → {_fmt_hyoji(name, v['当期'])}"
                f"（前期比 {v['前期比'] or '―'}）{v['フラグ']}")
        if v['注記']:
            line += f" ／ {v['注記']}"
        if v['確認事項']:
            line += f" ／ 確認事項: {v['確認事項']}"
        flagged_lines.append(line)
    flagged = '\n'.join(flagged_lines) if flagged_lines else '（フラグが立った指標はありません）'

    shokan = keisan['実質債務償還年数']
    if shokan is None and keisan.get('実質債務償還年数_算出不能理由'):
        shokan_hyoji = (
            f'{keisan["実質債務償還年数_算出不能理由"]}。'
            f'カルテのデータ欠損ではないので「データ欠損」と書かないこと。'
            f'この場合、縦軸5・6（償還年数による区分）の判定は行わない'
        )
    elif shokan is None:
        zeibiki = zaimu_cur.get('税引後利益')
        genka = zaimu_cur.get('減価償却費')
        if zeibiki is not None and genka is not None:
            shokan_hyoji = (
                f'算出不能（税引後利益{zeibiki}＋減価償却費{genka}＝{zeibiki + genka}で分母が0以下のため）。'
                f'カルテのデータ欠損ではないので「データ欠損」と書かないこと。'
                f'この場合、縦軸5・6（償還年数による区分）の判定は行わない'
            )
        else:
            shokan_hyoji = '算出不能（計算に必要な数値がカルテから取得できなかったため）'
    else:
        shokan_hyoji = f'{shokan}年'

    return f"""あなたは債務者区分・業況分析アシスタントの推論担当です。以下の機械判定結果をもとに、
特例適用の判断・会社概要要約・業界特異性所見・判定根拠の文章化のみを行ってください。
縦軸・横軸・交点区分そのものは既に機械判定済みであり、あなたが変更することはできません。

【会社概要（抽出データ）】※会社概要要約は、下記の5要素をすべて使って作成すること
①業種: {gaiyo.get('業種')}
　取扱品目・事業内容: {' / '.join(str(x) for x in (gaiyo.get('取扱品目') or [])) or '（記載なし）'}
②設立年: {gaiyo.get('設立年')}
　沿革: {' / '.join(str(x) for x in (gaiyo.get('沿革') or [])) or '（記載なし）'}
③資本金: {_num(gaiyo.get('資本金'))}百万円　上場区分: {gaiyo.get('上場') or '（記載なし）'}
　大株主・資本構成: {' / '.join(str(x) for x in (gaiyo.get('大株主') or [])) or '（記載なし）'}
④代表者: {gaiyo.get('代表者')}（{_num(gaiyo.get('代表者年令'))}才）
　役員: {' / '.join(str(x) for x in (gaiyo.get('役員') or [])) or '（記載なし）'}
　本店所在地: {gaiyo.get('本店所在地') or '（記載なし）'}　従業員数: {_num(gaiyo.get('従業員数'))}名
　主要銀行: {gaiyo.get('主要銀行') or '（記載なし）'}
⑤直近売上規模: {_num(gaiyo.get('直近売上規模'))}百万円
（商号: {gaiyo.get('商号')}　特記事項: {gaiyo.get('特記事項') or 'なし'}）

【カルテ記載の定性情報】※業界特異性所見の材料。ここに書かれていないことを推測で足さないこと
決算内容・特殊要因説明: {' / '.join(str(x) for x in (gaiyo.get('決算内容・特殊要因説明') or []))[:1200] or '（記載なし）'}
外部懸念情報: {gaiyo.get('外部懸念情報') or '（記載なし）'}
前回の債務者区分: {(gaiyo.get('グループ情報') or {}).get('区分') or '（記載なし）'}（判定日: {(gaiyo.get('グループ情報') or {}).get('債務者区分判定日') or '―'}）

【機械判定結果】
縦軸: {hantei['採用行']}　{hantei['採用行名']}
　縦軸の判定内訳: {jujiku_uchiwake}
横軸: {hantei['横軸記号']}　{hantei['横軸名']}（{hantei['横軸備考']}）
交点区分: {kouten_str}
低位区分優先を適用したか: {hantei['低位区分優先を適用したか']}（該当行候補: {hantei['該当行候補']}）
実質債務償還年数: {shokan_hyoji}

{daitai_note}
{teii_note}

【債務者区分判定実施基準 注①〜⑫（原本どおり）】
{CHUKI_ZENBUN}

※このうち「特例」（代替区分の適用可否を判断するもの）は 注{TOKUREI_NO_BANGO} のみです。
　{NON_TOKUREI_SETSUMEI}

【業況分析でフラグが立った指標】
※印は末尾に意味を併記している。★【改善】は良い変化なので、悪化として書かないこと。
※数値は出力Excelの表と同じ書式で示している。所見や判定根拠に書くときはこの表記のまま使うこと。
{flagged}

上記を踏まえ、特例適用判断・最終形式区分・判定根拠・会社概要要約（150字±10）・業界特異性所見を作成してください。
推測で数値を補完しない、根拠のない断定をしない、というルールを厳守してください。

【判定根拠の書き方（厳守）】
- 「・」で始まる箇条書き3〜5点。各点に具体的な数値根拠を入れる
- **箇条書きの各行にタグ（【事実】【意見】等）を付けてはならない。**
  タグを付けるのは「業界特異性所見」だけです
- **信頼度表示（【社内基準に基づく形式判定】/【推測を含む形式判定】）は、
  文章全体のいちばん最後に1回だけ書くこと。** 箇条書きの各行の末尾に繰り返し付けてはならない

【業界特異性所見のタグについて（審査資料として重要）】
本シートは審査担当者が最終判定を行うための資料であり、「カルテに書いてある事実」と
「書き手の評価」を読み手が区別できることが必須です。次を厳守してください。
- 「留意が必要」「懸念される」「注視すべき」など評価・判断を含む文は、
  根拠がカルテにあっても「事実」ではなく**「意見」**とすること
- **1要素に事実と評価を混ぜてはならない。** 混ぜると全部「意見」になり、
  カルテに書かれている事実がどれか読み手に分からなくなる。
  事実の要素を先に、評価の要素を後に、別の要素として書くこと
- **「事実」の要素を1個以上必ず含めること。**
  カルテの数値に触れているのに「事実」が1つもない所見は不可
- ここで書くのはカルテ由来の内容だけです。外部の業界情報はここでは扱いません
  （別の処理で後から付け加えます）"""



_GAIBU_SHOKEN_SCHEMA = {
    'type': 'object',
    'properties': {
        '所見要素': {
            'type': 'array',
            'description': (
                '外部リサーチ由来の着眼点。配列の1要素が1文にあたる。1個以上。'
                '★タグ（【一般論】等）は文中に書かない。文の本文だけを「文」に入れる。'
                '★1要素に「資料の内容」と「評価」を混ぜてはならない。'
                '資料の内容を述べる要素を先に、評価の要素を後に、別の要素として書く。'
                '★「一般論」の要素を1個以上必ず含めること。'
            ),
            'items': {
                'type': 'object',
                'properties': {
                    '種類': {'type': 'string', 'enum': ['一般論', '推測', '意見'],
                            'description':
                                '一般論＝資料に書かれている内容をそのまま述べる。評価の語を含めない。'
                                '推測＝資料の内容から推定できるが確認が取れていない内容。'
                                '意見＝評価・判断・推奨を含む文'
                                '（「留意が必要」「懸念される」等の語を含む文は意見）。'},
                    '出典資料番号': {'type': 'integer',
                                'description': '下に渡す資料一覧の番号。'},
                    '文': {'type': 'string',
                          'description': 'タグを付けない、文の本文だけ。「。」で終える。'},
                },
                'required': ['種類', '出典資料番号', '文'],
            },
        },
    },
    'required': ['所見要素'],
}


def _build_gaibu_prompt(gyoshu: str, chiiki: str, gyokai_joho: dict) -> str:
    zenbun = gyokai_joho['zenbun']
    lines = [
        '【外部リサーチで取得した業界情報】',
        '※これは業種一般の傾向であり、この会社を調べたものではない。参考情報として扱う。',
        '※★以下は年と対象業種をコード側で確認済みの資料です。'
        '渡したものは使ってよい。使うかどうかの判断は不要。',
        '※出典（媒体名）と年を文中に示すこと。年は下に書いてある値をそのまま使い、'
        '推測で書かないこと。',
        '※資料の対象（業種・地域・企業規模）が当社と違うなら、その違いを添えること'
        '（例：「四国地区を対象とした調査によれば」「大手企業を対象とした調査によれば」）。',
        '※この会社固有の事実として書いてはならない。業界の傾向を当社の数値の'
        '原因として断定しないこと（当社の数値そのものはこの後の処理で別に扱う）。',
        '★下の資料に「当社業種の記述」として引用されている内容だけを事実として使うこと。'
        'それ以外の数値を、記憶や推測で書き足してはならない。'
        '引用が表の形（ラベルと数値が並ぶ行）であっても、そこに書かれている数値は'
        'そのまま使ってよい（コード側で対応関係を確認済み）。'
        '引用が文章の場合も、その文に書かれている数値をそのまま使ってよい。',
        '★渡した資料の説明をそのまま書いてはならない。'
        '「本資料は〜の調査である」「調査事項には〜が含まれる」「調査対象は上位〇社」'
        'のように、その資料が何を調べているかを述べただけの文は所見に載せない。'
        '書くのは業界がどうなっているか（数値・水準・方向）である。'
        '要点に数値や傾向が無い資料は、無理に使わず他の資料を使うこと。',
        '★以下の資料のうち少なくとも1件に触れること。',
        '※複数渡している場合は複数に触れてよい。触れた資料は媒体名で特定できるように書くこと。',
        '※「何本の検索で出てきたか」は、その資料がその業界の一般的な情報源かどうかの目安。'
        '本数が多いものを優先して使う。所見に本数を書く必要はない。'
        '★ただし「毎回必ず取得している定点資料」は検索を通していないため本数を持たない。'
        '本数が無いことは優先度が低いことを意味しない。',
        '※年が「本文に記載がない」となっている資料は、年が新しいものより後に置いている。'
        '使う場合は「年次未掲載」と文中に示し、鮮度が確かめられないことが読み手に伝わるようにすること。',
        f"検索クエリ: {gyokai_joho.get('query')}",
        '',
    ]
    for z in zenbun:
        lines.append(f"{z['no']}. {z['title']}")
        lines.append(f"   出典URL: {z['url']}")
        if z.get('nen'):
            lines.append(f"   年: 発行{z['発行年']}／データの年次{z['データの年次']}")
        else:
            lines.append(f'   年: 本文に記載がない（{NEN_MIKEISAI}）')
            lines.append('   ★この資料を使う場合は、文中に「年次未掲載」と明記すること。'
                         '年を推測して書いてはならない')
        lines.append(f"   対象: 業種={z['対象業種']}／地域={z['対象地域']}"
                     f"／規模={z['対象企業規模']}（当社との一致: {z['業種の一致']}）")
        if z.get('固定'):
            lines.append('   この資料は毎回必ず取得している定点資料')
        else:
            lines.append(f"   この資料は{len(gyokai_joho.get('queries') or [1])}本の検索のうち"
                         f"{z.get('出現本数', 1)}本で出てきた"
                         + ('（一覧ページから深掘りして取得）' if z.get('深掘り') else ''))
        lines.append(f"   要点: {z['要点']}")

    return f"""あなたは債務者区分・業況分析アシスタントの推論担当です。
以下の外部リサーチの資料だけを材料に、業界一般の傾向を文章にしてください。
この会社固有の数値・事実は渡していません（別の処理で扱うため）。ここでは業界の話だけを書いてください。

【当社の業種】{gyoshu}
【当社の所在地】{chiiki}

{chr(10).join(lines)}
"""


async def _call_llm_gaibu_shoken(gyoshu: str, chiiki: str, gyokai_joho: dict,
                                 model: str = '') -> tuple[list, dict]:
    """外部リサーチ由来の所見要素だけを、カルテを見せない別ターンで生成する。

    メインの判定ターンと1つのllm_callにまとめず分けているのは、同じプロンプトに
    カルテの数値と外部資料の両方が入っていると、外部情報がカルテ由来として書かれる
    （出典を偽る）ことをコード側の検査に頼らず防ぐため。この処理はカルテのデータを
    一切受け取らないので、外部情報がカルテ由来として混ざることが構造的に起きない。

    zenbun が空なら呼ばない（呼び出し側で判定する）。
    """
    from agent_sdk import llm_call, ToolCallError

    prompt = _build_gaibu_prompt(gyoshu, chiiki, gyokai_joho)
    use_model = model or LLM_MODEL
    try:
        res = await llm_call(prompt=prompt, schema=_GAIBU_SHOKEN_SCHEMA, model=use_model)
    except ToolCallError as e:
        _glog(f'外部情報の所見生成に失敗しました（この情報なしで続行します）: {e}')
        return [], {'in': 0, 'out': 0}
    except Exception as e:
        _glog(f'外部情報の所見生成に失敗しました（この情報なしで続行します）: {type(e).__name__}: {e}')
        return [], {'in': 0, 'out': 0}

    usage = res.get('usage') or {}
    return (res.get('data') or {}).get('所見要素') or [], \
        {'in': usage.get('input_tokens') or 0, 'out': usage.get('output_tokens') or 0}

async def _call_llm_reasoning(extracted: dict, gyokyo_bunseki: dict[str, dict],
                              model: str = '') -> tuple[dict, str]:
    from agent_sdk import llm_call, ToolCallError

    prompt = _build_llm_prompt(extracted, gyokyo_bunseki)
    use_model = model or LLM_MODEL
    try:
        res = await llm_call(prompt=prompt, schema=_LLM_RESPONSE_SCHEMA, model=use_model)
    except ToolCallError as e:
        raise Exception(f'llm_call呼び出しに失敗しました（model={use_model}）: {e}') from e

    used_model = res.get('model', '(モデル名を取得できず)')
    usage = res.get('usage', {})
    print(f'[llm_call] model={used_model} '
          f'input_tokens={usage.get("input_tokens", "?")} '
          f'output_tokens={usage.get("output_tokens", "?")} '
          f'prompt_chars={len(prompt)}')
    logger.info(f'llm_call 実行: model={used_model} usage={usage}')

    return res['data'], used_model


OUTPUT_DIR = 'output'
OUTPUT_BASENAME = '債務者区分業況分析シート'
_FILENAME_NG = re.compile(r'[\\/:*?"<>|\r\n\t]')


def _output_path(shogo: str | None, output_path: str = '') -> str:
    name = _FILENAME_NG.sub('_', str(shogo or '判定').strip()) or '判定'
    d = os.path.dirname(output_path or '') or OUTPUT_DIR
    if d:
        os.makedirs(d, exist_ok=True)
    return os.path.join(d, f'{OUTPUT_BASENAME}_{name}.xlsx')


def _kouten_kara_saishu(kouten: str) -> str:
    m = re.match(r'^(.+?)(\d*)$', kouten)
    return f'{m.group(1)}先{m.group(2)}'


def _saishu_kensa(kouten: str, saishu: str) -> str:
    if '（' in kouten:
        return saishu
    tadashii = _kouten_kara_saishu(kouten)
    if saishu != tadashii:
        _glog(f'最終形式区分を直しました: {saishu!r} → {tadashii!r}'
              f'（交点「{kouten}」には特例適用の余地が無いため機械的に決まる）')
        return tadashii
    return saishu


async def process_karte(karte_path: str, template_path: str, output_path: str,
                        model: str = '') -> dict:
    extracted = await extract_and_calculate(karte_path)
    if extracted['status'] != 'ok':
        logger.info(f"判定を行わず終了（情報不足）: 不足項目={extracted.get('missing')}")
        print(f"[判定せず終了] 情報不足のため判定していません。不足項目: {extracted.get('missing')}")
        return extracted

    gyokyo_bunseki = _calculate_gyokyo_bunseki(extracted['zaimu'],
                                              extracted.get('karte_shihyo'))

    gyokai_joho = await search_gyokai_joho(
        extracted['gaiyo'].get('業種'), extracted['gaiyo'].get('取扱品目'))
    if ZENBUN_HANTEI:
        gyokai_joho = await zenbun_hantei(
            gyokai_joho,
            extracted['gaiyo'].get('業種') or '不明',
            str(extracted['gaiyo'].get('本店所在地') or '不明'))

    llm_result, used_model = await _call_llm_reasoning(
        extracted, gyokyo_bunseki, model=model)

    zenbun = (gyokai_joho or {}).get('zenbun') or []
    gaibu_items, gaibu_kosuto = ([], {'in': 0, 'out': 0})
    if zenbun:
        gaibu_items, gaibu_kosuto = await _call_llm_gaibu_shoken(
            extracted['gaiyo'].get('業種') or '不明',
            str(extracted['gaiyo'].get('本店所在地') or '不明'),
            gyokai_joho, model=model)
    if _GAIBU_TRACE:
        _GAIBU_TRACE.setdefault('コスト', {})
        _GAIBU_TRACE['コスト']['所見生成_外部ターン_in'] = gaibu_kosuto['in']
        _GAIBU_TRACE['コスト']['所見生成_外部ターン_out'] = gaibu_kosuto['out']

    gyokai_shoken, gyokai_no_inyou, mondai, kaz_kensa = _shoken_kumitate(
        llm_result['業界特異性所見'], gaibu_items, zenbun)
    if mondai:
        _glog(f'所見に不備のある文 {len(mondai)}件。該当文だけ書き直させます')
        for x in mondai:
            _glog(f'  {x["理由"]}｜{x["文"].strip()[:56]}')
        gyokai_shoken, shusei_kiroku = await _shoken_shusei(gyokai_shoken, mondai, model)
    else:
        shusei_kiroku = {'書き直しを依頼した文': 0, '書き直せた文': 0, '削除した文': 0}
    if _GAIBU_TRACE:
        _GAIBU_TRACE.setdefault('所見の検査', {}).update({**kaz_kensa, **shusei_kiroku})
        _GAIBU_TRACE['所見の検査']['不備の内訳'] = [
            {'理由': x['理由'], '文': x['文'].strip()[:80]} for x in mondai]


    kesson = bool(extracted.get('missing_items')) or any(
        v.get('フラグ') == FLAG_KESSON for v in gyokyo_bunseki.values())
    hantei_konkyo = _normalize_hantei_konkyo(
        llm_result['判定根拠'], extracted['hantei']['交点区分'], kesson)
    keikoku = _check_gaiyo_mojisu(llm_result['会社概要要約'])

    kakutei_path = _output_path(extracted['gaiyo'].get('商号'), output_path)
    if output_path and os.path.basename(output_path) and \
            os.path.basename(output_path) != os.path.basename(kakutei_path):
        logger.info(f'出力ファイル名を商号から確定しました: '
                    f'{os.path.basename(output_path)} → {os.path.basename(kakutei_path)}')

    result = await write_output_excel(
        template_path=template_path,
        output_path=kakutei_path,
        extracted=extracted,
        llm_tokurei_hanteibun=llm_result['特例適用判断'],
        llm_saishu_kubun=_saishu_kensa(extracted['hantei']['交点区分'], llm_result['最終形式区分']),
        llm_hantei_konkyo=hantei_konkyo,
        llm_kaisha_gaiyo_yoyaku=llm_result['会社概要要約'],
        gyokyo_bunseki=gyokyo_bunseki,
        llm_gyokai_shoken=gyokai_shoken,
        gyokai_no_inyou=gyokai_no_inyou,
        used_model=used_model,
        gyokai_joho=gyokai_joho,
    )
    gaibu_log_path = _gaibu_json(gyokai_joho, extracted['gaiyo'].get('商号'),
                                 os.path.dirname(output_path or '') or OUTPUT_DIR)
    result['status'] = 'ok'
    if keikoku:
        result['警告'] = [keikoku]
    if gaibu_log_path:
        result['外部リサーチログ'] = gaibu_log_path
    return result
